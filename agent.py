"""
feling. Sourcing Agent
======================
Monitors Grailed (US / Asia / Europe), Vestiaire, Yahoo Japan, and Yahoo Taiwan
for matching vintage pieces — multi-region lanes for price-competitive finds.
Sends findings to Telegram for approval.

Setup:
  1. Copy .env.example to .env and fill in values
  2. pip install -r requirements.txt
  3. python agent.py

Architecture:
  - Scheduler polls each platform every N minutes
  - New matches are deduplicated via seen_ids.json
  - Each match is sent to Telegram with photo + details + Approve/Skip buttons
  - Approve opens the purchase URL; Skip dismisses
"""

import os
import json
import logging
import asyncio
import hashlib
import re
import secrets
import smtplib
import subprocess
import html as html_lib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from io import BytesIO
from urllib.parse import quote, urljoin
from datetime import datetime, timedelta
from pathlib import Path
import bcrypt
from aiohttp import web
from dotenv import load_dotenv
from bs4 import BeautifulSoup
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from curl_cffi.requests import AsyncSession
from telegram import (
    Update, InlineKeyboardButton, InlineKeyboardMarkup, BotCommand,
    InputFile,
)
from telegram.ext import (
    Application, CommandHandler, CallbackQueryHandler, ContextTypes,
)

load_dotenv(override=True)
logging.basicConfig(
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    level=logging.INFO
)
log = logging.getLogger("feling")

# ─── Config ───────────────────────────────────────────────────────────────────

TELEGRAM_TOKEN   = os.getenv("TELEGRAM_TOKEN")
TELEGRAM_CHAT_ID = os.getenv("TELEGRAM_CHAT_ID")   # your personal chat id
POLL_INTERVAL    = int(os.getenv("POLL_INTERVAL_MINUTES", "30"))
REPORT_EMAIL     = (os.getenv("REPORT_EMAIL") or "").strip()
REPORT_FROM      = (os.getenv("REPORT_FROM") or os.getenv("SMTP_USER") or "").strip()
SMTP_HOST        = (os.getenv("SMTP_HOST") or "smtp.gmail.com").strip()
SMTP_PORT        = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER        = (os.getenv("SMTP_USER") or "").strip()
SMTP_PASSWORD    = (os.getenv("SMTP_PASSWORD") or "").strip()
WEEKLY_REPORT_DAY = (os.getenv("WEEKLY_REPORT_DAY") or "mon").strip().lower()
WEEKLY_REPORT_HOUR = int(os.getenv("WEEKLY_REPORT_HOUR", "9"))
PROCUREMENT_EXPORT_DAY  = (os.getenv("PROCUREMENT_EXPORT_DAY") or "fri").strip().lower()
PROCUREMENT_EXPORT_HOUR = int(os.getenv("PROCUREMENT_EXPORT_HOUR", "9"))

# Team web API (login required) — see /ops, /review. Railway (and most PaaS
# hosts) inject PORT and expect the app to bind to it; WEB_PORT is the local
# dev fallback.
WEB_PORT = int(os.getenv("PORT") or os.getenv("WEB_PORT", "8787"))

# All state lives under DATA_DIR so it survives redeploys when it points at a
# mounted volume (e.g. Railway: set DATA_DIR=/data and attach a volume there).
# Defaults to the repo checkout for local dev, matching prior behavior exactly.
DATA_DIR = Path(os.getenv("DATA_DIR", "."))
DATA_DIR.mkdir(parents=True, exist_ok=True)

SEEN_DB          = DATA_DIR / "seen_ids.json"
PENDING_DB       = DATA_DIR / "pending.json"
APPROVED_DB      = DATA_DIR / "approved.json"
SETTINGS_DB      = DATA_DIR / "settings.json"
TREND_DB         = DATA_DIR / "trend_snapshot.json"
TASTE_DB         = DATA_DIR / "taste_profile.json"
SKIPPED_DB       = DATA_DIR / "skipped.json"
EXPORT_MARKER_DB = DATA_DIR / "last_procurement_export.json"
INVENTORY_STATUS_DB = DATA_DIR / "inventory_status.json"  # written by web/ /ops page
USERS_DB = DATA_DIR / "team_users.json"  # named team accounts — see `python agent.py --add-user`
SESSION_TTL_HOURS = 24 * 7  # web login session lifetime
REPORTS_DIR      = DATA_DIR / "reports"


def _seed_data_dir_from_repo():
    """One-time migration for a fresh volume: if DATA_DIR is separate from the
    repo checkout (e.g. a brand-new Railway volume) and doesn't have a given
    file yet, but the repo checkout does, copy it over so history already
    committed to git isn't lost the first time persistent storage is wired up.
    """
    if DATA_DIR.resolve() == Path(".").resolve():
        return
    for filename in (
        "seen_ids.json", "pending.json", "approved.json", "settings.json",
        "trend_snapshot.json", "taste_profile.json", "skipped.json",
        "last_procurement_export.json", "inventory_status.json", "team_users.json",
    ):
        dest = DATA_DIR / filename
        src = Path(filename)
        if not dest.exists() and src.exists():
            dest.write_text(src.read_text(encoding="utf-8"), encoding="utf-8")
            log.info("Seeded %s into %s from repo checkout", filename, DATA_DIR)


_seed_data_dir_from_repo()
HTTP_TIMEOUT_SECONDS = 30
GRAILED_ALGOLIA_APP_ID = "MNRWEFSS2Q"
GRAILED_ALGOLIA_SEARCH_KEY = "c89dbaddf15fe70e1941a109bf7c2a3d"
GRAILED_LISTING_INDEX = "Listing_by_date_added_production"
GRAILED_HEAT_INDEX = "Listing_by_heat_production"
VESTIAIRE_SITEMAP_INDEX = "https://www.vestiairecollective.com/sitemaps/https_sitemap-en.xml"
VESTIAIRE_PRODUCT_SITEMAP_LIMIT = 10
VESTIAIRE_MATCHING_URL_LIMIT = 20
VESTIAIRE_SITEMAP_FETCH_CONCURRENCY = 6
VESTIAIRE_PRODUCT_FETCH_CONCURRENCY = 8
TASTE_MIN_SAMPLES = 3
WEEKLY_REPORT_LOOKBACK_DAYS = 7
WEEKLY_REPORT_MAX_ITEMS = 40
PROCUREMENT_EXPORT_LOOKBACK_DAYS = 7  # fallback window before any export has ever run

# Mirrors web/src/types.ts STATUS_LABEL — keep both in sync.
STATUS_LABELS = {
    "approved": "Approved (not yet purchased)",
    "purchased": "Purchased (in transit)",
    "received": "Received",
    "qc_passed": "QC passed",
    "photographed": "Photographed",
    "listed": "Listed for sale",
    "sold": "Sold",
    "returned": "Returned / issue",
}
_TASTE_CACHE: dict | None = None
_BOT_REF = None  # set once in main() so shared approve/skip logic can notify Telegram
_SESSIONS: dict[str, dict] = {}  # in-memory web session store: token -> {username, expires}

CONDITION_RANK = {
    "is_new": 5,
    "excellent": 5,
    "new": 5,
    "very good": 4,
    "gently_used": 4,
    "is_gently_used": 4,
    "good": 3,
    "is_used": 2,
    "used": 2,
    "a": 5,
    "b": 3,
    "c": 2,
    "fair": 1,
    "unknown": 0,
}

TREND_SIGNALS = [
    # Chanel / Dior rare bags
    "chanel", "2.55", "classic flap", "jumbo", "kelly", "saddle", "trotter",
    "columbus", "rare", "archive", "limited",
    # Dior shoes + clothes
    "dior", "galliano", "bondage", "harness", "heel", "pump", "oblique",
    # Cavalli / Gaultier / Tom Ford Gucci
    "cavalli", "leopard", "python", "animal print", "gaultier", "cone bra",
    "tom ford", "gucci", "jackie", "bamboo", "horsebit",
    "corset", "y2k", "2000s", "vintage", "runway", "slay",
]

STOP_WORDS = {
    "the", "and", "for", "with", "size", "new", "nwt", "womens", "women",
    "woman", "dress", "item", "from", "this", "that", "very", "good",
    "condition", "used", "gently", "black", "white", "red", "blue",
}

# ─── Sourcing Rules ───────────────────────────────────────────────────────────
# Aesthetic north stars: Isle of Monday, Jean Vintage, Break Archive
# Chanel rare bags • Dior rare bags / shoes / clothes
# Cavalli • Galliano • Tom Ford Gucci • Jean Paul Gaultier

CLOTHES_INCLUDE = [
    "dress", "mini dress", "slip dress", "gown", "skirt", "mini skirt",
    "corset", "bustier", "camisole", "halter", "bodysuit", "top", "blouse",
    "set", "co-ord", "coord", "two piece", "lace", "mesh", "sheer",
    "silk", "satin", "ruched", "cutout", "bodycon", "low rise",
    "jacket", "coat",
]

BAGS_INCLUDE = [
    "bag", "handbag", "clutch", "shoulder bag", "tote", "purse", "pochette",
    "saddle", "crossbody", "flap", "2.55", "jackie",
]

SHOES_INCLUDE = [
    "heels", "heel", "sandal", "mule", "stiletto", "pump", "pumps",
    "boot", "boots", "shoe", "shoes",
]

BAGS_SHOES_INCLUDE = BAGS_INCLUDE + SHOES_INCLUDE

ARCHIVE_SLAY_INCLUDE = [
    "galliano", "john galliano", "dior by galliano", "columbus",
    "bondage", "harness", "strappy", "cage", "fetish", "corsetry",
    "newspaper print", "saddle", "trotter", "oblique", "monogram",
    "tom ford", "jackie", "bamboo", "horsebit",
    "chanel", "2.55", "classic flap", "quilted", "cc", "double flap",
    "rare", "limited", "collector", "archive", "runway",
    "leopard", "python", "animal print", "tiger", "zebra", "snake",
    "tattoo", "baroque", "printed", "embellished", "beaded", "sequined",
    "crystal", "rhinestone", "vintage", "y2k", "2000s", "90s",
    "fw", "ss", "gaultier", "jean paul gaultier", "cone", "tattoo mesh",
]

STYLE_EXCLUDE = [
    "plain", "minimalist", "basic", "simple", "office", "school", "uniform",
    "sweatshirt", "hoodie", "sneaker", "trainer", "wallet",
    "menswear", "men's", " for men", "football", "jersey",
]

CONDITIONS = [
    "Very Good", "Excellent", "gently_used", "is_gently_used",
    "is_new", "is_used", "used", "A", "B",
]

JP_FEMININE_INCLUDE = [
    "ワンピース", "ドレス", "スカート", "バッグ", "ハンドバッグ", "クラッチ",
    "ヒール", "パンプス", "ミュール", "サンダル", "トップス", "ブラウス",
    "ビスチェ", "コルセット", "キャミ", "ボディスーツ", "セットアップ",
    "ジャケット", "コート", "レース", "シルク", "サテン", "靴",
]

# Traditional (TW) + Simplified — Yahoo TW / Grailed Asia / mixed titles
CN_FEMININE_INCLUDE = [
    "洋裝", "洋装", "連身裙", "连身裙", "裙子", "上衣", "外套",
    "包", "手袋", "手拿包", "肩背包", "馬鞍包", "马鞍包",
    "高跟鞋", "跟鞋", "穆勒鞋", "涼鞋", "凉鞋", "靴", "鞋",
    "胸衣", "緊身胸衣", "紧身胸衣", "馬甲", "马甲", "網紗", "网纱",
    "套裝", "套装",
]

CJK_SHAPE_INCLUDE = JP_FEMININE_INCLUDE + CN_FEMININE_INCLUDE

# Map native script → English tokens so JP/CN titles still hit taste + style gates
CJK_ALIASES = {
    "シャネル": ["chanel"],
    "香奈兒": ["chanel"],
    "香奈尔": ["chanel"],
    "ディオール": ["dior"],
    "クリスチャンディオール": ["dior", "christian dior"],
    "迪奧": ["dior"],
    "迪奥": ["dior"],
    "ガリアーノ": ["galliano", "john galliano"],
    "ジョンガリアーノ": ["galliano", "john galliano"],
    "加利亞諾": ["galliano"],
    "加利亚诺": ["galliano"],
    "カヴァリ": ["cavalli"],
    "カバリ": ["cavalli"],
    "ロベルトカヴァリ": ["cavalli", "roberto cavalli"],
    "ジャストカヴァリ": ["cavalli", "just cavalli"],
    "卡沃利": ["cavalli"],
    "卡瓦利": ["cavalli"],
    "ゴルチエ": ["gaultier"],
    "ジャンポールゴルチエ": ["gaultier", "jean paul gaultier"],
    "ジャンポール・ゴルチエ": ["gaultier", "jean paul gaultier"],
    "高緹耶": ["gaultier"],
    "高缇耶": ["gaultier"],
    "高提耶": ["gaultier"],
    "グッチ": ["gucci"],
    "古馳": ["gucci"],
    "古驰": ["gucci"],
    "トムフォード": ["tom ford"],
    "トム・フォード": ["tom ford"],
    "湯姆福特": ["tom ford"],
    "汤姆福特": ["tom ford"],
    "サドル": ["saddle"],
    "トロッター": ["trotter"],
    "コロンブス": ["columbus"],
    "ジャッキー": ["jackie"],
    "バンブー": ["bamboo"],
    "ヴィンテージ": ["vintage"],
    "ビンテージ": ["vintage"],
    "アーカイブ": ["archive"],
    "レオパード": ["leopard"],
    "アニマル": ["animal print"],
    "コルセット": ["corset"],
    "ワンピース": ["dress"],
    "ドレス": ["dress"],
    "メッシュ": ["mesh"],
    "パンプス": ["pump", "heel", "shoe"],
    "ミュール": ["mule", "shoe"],
    "ヒール": ["heel", "shoe"],
    "靴": ["shoe"],
    "バッグ": ["bag"],
    "手袋": ["bag", "handbag"],
    "手拿包": ["clutch", "bag"],
    "手提包": ["bag", "handbag"],
    "肩背包": ["bag"],
    "包": ["bag"],
    "洋裝": ["dress"],
    "洋装": ["dress"],
    "連身裙": ["dress"],
    "连身裙": ["dress"],
    "馬鞍包": ["saddle", "bag"],
    "马鞍包": ["saddle", "bag"],
    "古董": ["vintage"],
    "稀有": ["rare"],
    "豹紋": ["leopard"],
    "豹纹": ["leopard"],
    "高跟鞋": ["heel", "shoe"],
    "胸衣": ["corset"],
}

JPY_PER_USD = 150
TWD_PER_USD = 32
DEFAULT_JAPAN_MAX_USD = 450
DEFAULT_TAIWAN_MAX_USD = 500

# Grailed location facets used for price-competitive regional sourcing
GRAILED_REGIONS = (
    ("United States", "US"),
    ("Asia", "Asia"),       # JP / TW / Asia sellers
    ("Europe", "Europe"),   # Italy + EU bargain lane
)

RULES = [
    {
        "brand": "Chanel Rare Bags",
        "focus": "bags",
        "keywords": [
            "chanel", "シャネル", "香奈兒", "香奈尔",
        ],
        "search_queries": [
            "chanel classic flap vintage",
            "chanel 2.55 vintage",
            "chanel jumbo rare",
            "chanel rare handbag archive",
            "chanel limited edition bag",
            "chanel vintage quilted bag",
            "chanel collector bag",
        ],
        "jp_queries": [
            "シャネル バッグ",
            "シャネル ヴィンテージ バッグ",
            "シャネル 2.55",
            "シャネル クラシックフラップ",
            "シャネル マトラッセ",
            "シャネル レア バッグ",
        ],
        "tw_queries": [
            "香奈兒 包",
            "香奈兒 古董 包",
            "香奈兒 2.55",
            "香奈兒 流浪包",
            "香奈兒 稀有 包",
        ],
        "cn_queries": [
            "香奈尔 包",
            "香奈尔 古董",
            "香奈尔 2.55",
        ],
        "max_price": 3500,
        "min_price": 200,
        "jp_min_price": 80,
        "jp_max_price": 1200,
        "conditions": CONDITIONS,
        "include": BAGS_INCLUDE + ["バッグ", "ハンドバッグ", "クラッチ"],
        "style": ARCHIVE_SLAY_INCLUDE + [
            "シャネル", "キルティング", "レア", "アーカイブ", "限定",
        ],
        "exclude": STYLE_EXCLUDE + ["sneaker", "wallet", "sunglass", "perfume", "scarf"],
        "era": [
            "rare", "archive", "limited", "vintage", "collector", "runway",
            "2.55", "classic flap", "jumbo", "90s", "2000s", "y2k",
            "レア", "アーカイブ", "限定", "ヴィンテージ",
        ],
    },
    {
        "brand": "Dior Rare Bags",
        "focus": "bags",
        "keywords": [
            "dior", "christian dior", "ディオール", "サドル", "トロッター",
            "迪奧", "迪奥", "馬鞍", "马鞍",
        ],
        "search_queries": [
            "dior saddle bag rare",
            "dior trotter bag archive",
            "dior columbus bag",
            "dior galliano saddle",
            "dior limited edition bag",
            "dior vintage monogram bag",
            "christian dior rare handbag",
        ],
        "jp_queries": [
            "ディオール サドル",
            "ディオール サドルバッグ",
            "ディオール トロッター",
            "ディオール コロンブス",
            "ガリアーノ ディオール バッグ",
            "ディオール バッグ ヴィンテージ",
        ],
        "tw_queries": [
            "迪奧 馬鞍包",
            "迪奧 包 古董",
            "迪奧 加利亞諾 包",
            "Dior saddle",
            "迪奧 手提包",
        ],
        "cn_queries": [
            "迪奥 马鞍包",
            "迪奥 包 古董",
            "迪奥 Galliano 包",
        ],
        "max_price": 2500,
        "min_price": 80,
        "jp_min_price": 30,
        "jp_max_price": 800,
        "conditions": CONDITIONS,
        "include": BAGS_INCLUDE + ["バッグ", "ハンドバッグ", "クラッチ"],
        "style": ARCHIVE_SLAY_INCLUDE + [
            "サドル", "トロッター", "ガリアーノ", "ディオール", "レア", "アーカイブ",
        ],
        "exclude": STYLE_EXCLUDE + ["sneaker", "wallet", "perfume"],
        "era": [
            "saddle", "trotter", "columbus", "galliano", "rare", "archive",
            "limited", "vintage", "1997", "1998", "1999", "2000", "2001",
            "2002", "2003", "2004", "2005", "90s", "2000s", "y2k",
            "サドル", "トロッター", "レア",
        ],
    },
    {
        "brand": "Dior Shoes",
        "focus": "shoes",
        "keywords": [
            "dior", "christian dior", "galliano", "ディオール", "ガリアーノ",
            "迪奧", "迪奥", "加利亞諾", "加利亚诺",
        ],
        "search_queries": [
            "dior galliano heels",
            "dior bondage heel",
            "dior vintage heels",
            "dior harness heel",
            "john galliano dior shoe",
            "dior stiletto archive",
            "dior mule vintage",
        ],
        "jp_queries": [
            "ディオール パンプス",
            "ディオール ヒール",
            "ディオール ガリアーノ 靴",
            "ガリアーノ期 パンプス",
            "ディオール ミュール",
            "ディオール ボンデージ",
            "ディオール サンダル",
        ],
        "tw_queries": [
            "迪奧 高跟鞋",
            "迪奧 加利亞諾 鞋",
            "迪奧 古董 鞋",
            "Dior Galliano heels",
            "迪奧 穆勒鞋",
        ],
        "cn_queries": [
            "迪奥 高跟鞋",
            "迪奥 加利亚诺 鞋",
            "迪奥 鞋 古董",
        ],
        "max_price": 900,
        "min_price": 50,
        "jp_min_price": 20,
        "jp_max_price": 400,
        "conditions": CONDITIONS,
        "include": SHOES_INCLUDE + ["ヒール", "パンプス", "ミュール", "サンダル", "靴"],
        "style": ARCHIVE_SLAY_INCLUDE + ["ガリアーノ", "ディオール", "アーカイブ"],
        "exclude": STYLE_EXCLUDE + ["sneaker", "trainer", "wallet"],
        "era": [
            "galliano", "bondage", "harness", "archive", "vintage",
            "1998", "1999", "2000", "2001", "2002", "2003", "2004",
            "90s", "2000s", "y2k", "ガリアーノ",
        ],
    },
    {
        "brand": "Dior Clothes",
        "focus": "clothes",
        "keywords": [
            "dior", "christian dior", "galliano", "ディオール", "ガリアーノ",
            "迪奧", "迪奥", "加利亞諾", "加利亚诺",
        ],
        "search_queries": [
            "dior galliano dress",
            "christian dior vintage dress",
            "dior by john galliano",
            "dior newspaper print",
            "dior columbus dress",
            "dior corset vintage",
            "dior runway dress y2k",
        ],
        "jp_queries": [
            "ディオール ガリアーノ ワンピース",
            "ディオール ドレス ヴィンテージ",
            "ガリアーノ ディオール",
            "ディオール コルセット",
            "ガリアーノ期 ワンピース",
            "ディオール ガリアーノ トップス",
        ],
        "tw_queries": [
            "迪奧 加利亞諾 洋裝",
            "迪奧 古董 洋裝",
            "Dior Galliano dress",
            "迪奧 胸衣",
        ],
        "cn_queries": [
            "迪奥 加利亚诺 洋装",
            "迪奥 连衣裙 古董",
            "迪奥 Galliano",
        ],
        "max_price": 1200,
        "min_price": 60,
        "jp_min_price": 25,
        "jp_max_price": 500,
        "conditions": CONDITIONS,
        "include": CLOTHES_INCLUDE + JP_FEMININE_INCLUDE,
        "style": ARCHIVE_SLAY_INCLUDE + ["ガリアーノ", "ディオール", "アーカイブ"],
        "exclude": STYLE_EXCLUDE + ["sneaker", "wallet", "bag", "heel"],
        "era": [
            "galliano", "archive", "vintage", "runway", "newspaper",
            "columbus", "1997", "1998", "1999", "2000", "2001", "2002",
            "2003", "2004", "2005", "90s", "2000s", "y2k", "ガリアーノ",
        ],
    },
    {
        "brand": "Roberto Cavalli",
        "focus": "clothes",
        "keywords": [
            "cavalli", "just cavalli", "roberto cavalli",
            "カヴァリ", "カバリ", "ロベルトカヴァリ", "ジャストカヴァリ",
            "卡沃利", "卡瓦利",
        ],
        "search_queries": [
            "roberto cavalli animal print dress",
            "cavalli leopard dress",
            "cavalli python dress",
            "roberto cavalli vintage dress",
            "cavalli corset",
            "cavalli silk set",
            "just cavalli y2k dress",
            "cavalli runway dress",
        ],
        "jp_queries": [
            "ロベルトカヴァリ ワンピース",
            "カヴァリ レオパード",
            "カヴァリ アニマル",
            "ジャストカヴァリ ドレス",
            "カヴァリ セットアップ",
            "ロベルトカヴァリ ドレス",
            "カヴァリ コルセット",
        ],
        "tw_queries": [
            "Cavalli 豹紋",
            "Cavalli 洋裝",
            "Roberto Cavalli 古董",
            "卡沃利 洋裝",
            "Cavalli 套裝",
        ],
        "cn_queries": [
            "卡沃利 连衣裙",
            "Cavalli 豹纹",
            "罗伯托卡沃利",
        ],
        "max_price": 700,
        "min_price": 40,
        "jp_min_price": 15,
        "jp_max_price": 350,
        "conditions": CONDITIONS,
        "include": CLOTHES_INCLUDE + JP_FEMININE_INCLUDE,
        "style": ARCHIVE_SLAY_INCLUDE + ["カヴァリ", "レオパード", "アニマル", "パイソン"],
        "exclude": STYLE_EXCLUDE + ["wallet", "sneaker"],
        "era": [
            "90s", "2000s", "00s", "y2k", "leopard", "python", "animal",
            "1999", "2000", "2001", "2002", "2003", "2004", "2005", "2006",
            "レオパード", "アニマル", "ランウェイ",
        ],
    },
    {
        "brand": "Galliano Clothes",
        "focus": "clothes",
        "keywords": [
            "galliano", "john galliano", "ガリアーノ", "ジョンガリアーノ",
            "加利亞諾", "加利亚诺",
        ],
        "search_queries": [
            "john galliano vintage dress",
            "galliano runway dress",
            "galliano corset",
            "galliano newspaper print",
            "john galliano silk dress",
            "galliano archive dress",
        ],
        "jp_queries": [
            "ジョンガリアーノ ワンピース",
            "ガリアーノ ドレス",
            "ガリアーノ アーカイブ",
            "ガリアーノ コルセット",
            "ジョンガリアーノ ヴィンテージ",
        ],
        "tw_queries": [
            "加利亞諾 洋裝",
            "Galliano 洋裝",
            "John Galliano 古董",
            "Galliano 胸衣",
        ],
        "cn_queries": [
            "加利亚诺 洋装",
            "John Galliano 连衣裙",
            "Galliano 古董",
        ],
        "max_price": 1100,
        "min_price": 50,
        "jp_min_price": 20,
        "jp_max_price": 450,
        "conditions": CONDITIONS,
        "include": CLOTHES_INCLUDE + JP_FEMININE_INCLUDE,
        "style": ARCHIVE_SLAY_INCLUDE + ["ガリアーノ", "アーカイブ"],
        "exclude": STYLE_EXCLUDE + ["sneaker", "wallet"],
        "era": [
            "archive", "vintage", "runway", "y2k", "2000s", "90s",
            "1997", "1998", "1999", "2000", "2001", "2002", "2003", "2004",
            "ガリアーノ", "アーカイブ",
        ],
    },
    {
        "brand": "Tom Ford Gucci",
        "focus": "clothes",
        "keywords": [
            "gucci", "tom ford", "グッチ", "トムフォード",
            "古馳", "古驰", "湯姆福特", "汤姆福特",
        ],
        "search_queries": [
            "tom ford gucci dress",
            "gucci by tom ford",
            "tom ford gucci vintage",
            "gucci jackie bag tom ford",
            "gucci bamboo tom ford",
            "tom ford gucci horsebit",
            "gucci tom ford y2k",
        ],
        "jp_queries": [
            "グッチ トムフォード",
            "トムフォード グッチ ドレス",
            "グッチ ジャッキー",
            "グッチ バンブー ヴィンテージ",
            "トムフォード グッチ ワンピース",
        ],
        "tw_queries": [
            "古馳 湯姆福特",
            "Gucci Tom Ford",
            "古馳 古董",
            "Gucci Jackie 包",
        ],
        "cn_queries": [
            "古驰 汤姆福特",
            "古驰 古董",
            "Gucci Tom Ford",
        ],
        "max_price": 1200,
        "min_price": 60,
        "jp_min_price": 25,
        "jp_max_price": 500,
        "conditions": CONDITIONS,
        "include": CLOTHES_INCLUDE + BAGS_SHOES_INCLUDE + JP_FEMININE_INCLUDE,
        "style": ARCHIVE_SLAY_INCLUDE + [
            "トムフォード", "グッチ", "ジャッキー", "バンブー", "アーカイブ",
        ],
        "exclude": STYLE_EXCLUDE + ["sneaker", "wallet", "belt only"],
        "era": [
            "tom ford", "jackie", "bamboo", "horsebit", "archive", "vintage",
            "1995", "1996", "1997", "1998", "1999", "2000", "2001", "2002",
            "2003", "2004", "90s", "2000s", "y2k", "トムフォード",
        ],
    },
    {
        "brand": "Jean Paul Gaultier",
        "focus": "clothes",
        "keywords": [
            "gaultier", "jean paul gaultier",
            "ゴルチエ", "ジャンポールゴルチエ",
            "高緹耶", "高缇耶", "高提耶",
        ],
        "search_queries": [
            "jean paul gaultier dress",
            "gaultier corset vintage",
            "gaultier mesh top",
            "jean paul gaultier tattoo",
            "gaultier conical bra",
            "jpg vintage dress",
            "gaultier archive dress",
        ],
        "jp_queries": [
            "ジャンポールゴルチエ ワンピース",
            "ゴルチエ コルセット",
            "ゴルチエ メッシュ",
            "ゴルチエ アーカイブ",
            "ゴルチエ ドレス",
            "ジャンポールゴルチエ ヴィンテージ",
        ],
        "tw_queries": [
            "高緹耶 洋裝",
            "高緹耶 胸衣",
            "JPG 洋裝",
            "Gaultier 古董",
            "高提耶 網紗",
        ],
        "cn_queries": [
            "高缇耶 连衣裙",
            "高缇耶 胸衣",
            "Gaultier 网纱",
        ],
        "max_price": 900,
        "min_price": 40,
        "jp_min_price": 15,
        "jp_max_price": 400,
        "conditions": CONDITIONS,
        "include": CLOTHES_INCLUDE + JP_FEMININE_INCLUDE,
        "style": ARCHIVE_SLAY_INCLUDE + ["ゴルチエ", "アーカイブ"],
        "exclude": STYLE_EXCLUDE + ["sneaker", "wallet"],
        "era": [
            "archive", "vintage", "corset", "tattoo", "mesh", "cone",
            "90s", "2000s", "y2k", "runway",
            "1995", "1996", "1997", "1998", "1999", "2000", "2001", "2002",
            "ゴルチエ", "アーカイブ",
        ],
    },
]

# ─── Helpers ──────────────────────────────────────────────────────────────────

def load_json(path: Path, default):
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            log.warning("Invalid JSON in %s, using default value.", path)
    return default

def save_json(path: Path, data):
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")

def make_id(url: str) -> str:
    return hashlib.md5(url.encode()).hexdigest()

def display_brand(brand: str | None) -> str:
    """Customer-facing brand label. Slay Outfit Archive → Others everywhere."""
    name = (brand or "").strip()
    if name.lower() == "slay outfit archive":
        return "Others"
    return name or "Others"

def normalize_item_brand(item: dict) -> dict:
    if "brand" in item:
        item = {**item, "brand": display_brand(item.get("brand"))}
    return item


def _canon_url(url: str | None) -> str:
    if not url:
        return ""
    return str(url).split("?", 1)[0].rstrip("/").lower()


def listing_uid(item: dict) -> str:
    return f"{item.get('platform')}_{item.get('id')}"


def upsert_approved(item: dict) -> list[dict]:
    """Append approved item while replacing same listing id or same URL."""
    approved = load_json(APPROVED_DB, [])
    uid = listing_uid(item)
    url = _canon_url(item.get("url"))
    kept: list[dict] = []
    for a in approved:
        if listing_uid(a) == uid:
            continue
        if url and _canon_url(a.get("url")) == url:
            continue
        kept.append(a)
    kept.append(item)
    return kept


_SHOE_RE = re.compile(
    r"(heel|heels|shoe|shoes|mule|mules|pump|pumps|boot|boots|sandal|sandals|"
    r"stiletto|sneaker|sneakers|loafer|loafers|espadrille|wedge|wedges|ballerin|slingback)",
    re.I,
)
_BAG_RE = re.compile(
    r"(handbag|clutch|purse|tote|cambon|columbus|trotter|crossbody|baguette|"
    r"boston|shoulder bag|saddle bag|\bsaddle\b|flap bag|\bbag\b)",
    re.I,
)


def categorize_item(item: dict) -> str:
    """Mirror web sync-catalog.mjs — title wins over sourcing lane brand."""
    title = str(item.get("title") or "")
    brand = re.sub(
        r"\brare bags?\b|\bbags?\s*&\s*shoes\b|\b(bags?|shoes?)\b",
        " ",
        str(item.get("brand") or ""),
        flags=re.I,
    )
    brand = re.sub(r"\s+", " ", brand).strip()
    if _SHOE_RE.search(title):
        return "shoes"
    if _BAG_RE.search(title):
        return "bags"
    hay = f"{brand} {title}"
    if _SHOE_RE.search(hay):
        return "shoes"
    if _BAG_RE.search(hay):
        return "bags"
    return "ready-to-wear"


def publish_to_website(item: dict | None = None) -> str:
    """Sync catalog; prep cutout when the approved piece is clothing or shoes."""
    web = Path("web")
    node = "node"
    sync = web / "scripts" / "publish-catalog.mjs"
    if not sync.exists():
        return "publish script missing"

    cat = categorize_item(item) if item else ""
    args = [node, str(sync)]
    if item and cat in ("ready-to-wear", "shoes", "bags"):
        args.append("--cutouts")
        args.append(f"--ids={listing_uid(item)}")
    else:
        # Always refresh catalog; cutouts catch up via npm run publish
        pass

    try:
        r = subprocess.run(
            args,
            cwd=Path("."),
            capture_output=True,
            text=True,
            timeout=180,
        )
        tail = (r.stdout or r.stderr or "").strip().splitlines()[-3:]
        status = "ok" if r.returncode == 0 else f"fail({r.returncode})"
        return status + ((" · " + " | ".join(tail)) if tail else "")
    except Exception as e:
        log.warning("Website publish failed: %s", e)
        return f"fail: {e}"


def load_settings() -> dict:
    return load_json(SETTINGS_DB, {
        "global_min_price": None,
        "global_max_price": None,
    })

def save_settings(settings: dict):
    save_json(SETTINGS_DB, settings)

def effective_price_bounds(rule: dict, platform: str | None = None) -> tuple[float, float]:
    """Apply optional global /price override; Asia auction lanes use cheaper bounds."""
    settings = load_settings()
    asia_lane = platform in {
        "YahooJP", "Buyee", "Japan", "YahooTW", "Taiwan",
    }
    if asia_lane:
        if platform in {"YahooTW", "Taiwan"}:
            min_price = float(rule.get("tw_min_price", rule.get("jp_min_price", 15)))
            max_price = float(
                rule.get("tw_max_price", rule.get("jp_max_price", DEFAULT_TAIWAN_MAX_USD))
            )
        else:
            min_price = float(rule.get("jp_min_price", 15))
            max_price = float(rule.get("jp_max_price", DEFAULT_JAPAN_MAX_USD))
    else:
        min_price = float(rule["min_price"])
        max_price = float(rule["max_price"])
    if settings.get("global_min_price") is not None:
        min_price = float(settings["global_min_price"])
    if settings.get("global_max_price") is not None:
        max_price = float(settings["global_max_price"])
    # Optional Japan-only ceiling via /price japan 300
    if platform in {"YahooJP", "Buyee", "Japan"} and settings.get("japan_max_price") is not None:
        max_price = min(max_price, float(settings["japan_max_price"]))
    if min_price > max_price:
        min_price, max_price = max_price, min_price
    return min_price, max_price


def expand_cjk_hay(text: str) -> str:
    """Append English aliases when a JP/CN brand or shape word is present."""
    extra: list[str] = []
    for native, latins in CJK_ALIASES.items():
        if native and native in text:
            extra.extend(latins)
    if not extra:
        return text
    return f"{text} {' '.join(dict.fromkeys(extra))}"


def cjk_search_queries(rule: dict) -> list[str]:
    return list(dict.fromkeys([
        *(rule.get("jp_queries") or []),
        *(rule.get("tw_queries") or []),
        *(rule.get("cn_queries") or []),
    ]))


def extract_taste_tokens(text: str) -> list[str]:
    """Pull aesthetic tokens from a listing title/brand string."""
    hay = expand_cjk_hay(text.lower())
    found = []
    vocab = list(dict.fromkeys(
        TREND_SIGNALS + ARCHIVE_SLAY_INCLUDE + CLOTHES_INCLUDE + BAGS_SHOES_INCLUDE
        + list(CJK_ALIASES.keys())
    ))
    for term in vocab:
        if term in hay:
            found.append(term)
    # Extra free tokens from title words
    for word in re.findall(r"[a-z0-9']{3,}", hay):
        if word in STOP_WORDS:
            continue
        if word.isdigit():
            continue
        if word not in found:
            found.append(word)
    return found


def rebuild_taste_profile() -> dict:
    """
    Retrain aesthetic preferences from Approve / Skip feedback.
    This is lightweight preference learning, not a neural model.
    """
    approved = load_json(APPROVED_DB, [])
    skipped = load_json(SKIPPED_DB, [])

    pos_counts: dict[str, int] = {}
    neg_counts: dict[str, int] = {}
    brand_counts: dict[str, int] = {}
    prices = []

    for item in approved:
        brand = item.get("brand", "unknown")
        brand_counts[brand] = brand_counts.get(brand, 0) + 1
        if item.get("price") is not None:
            prices.append(float(item["price"]))
        for token in extract_taste_tokens(f"{brand} {item.get('title', '')}"):
            pos_counts[token] = pos_counts.get(token, 0) + 1

    for item in skipped:
        for token in extract_taste_tokens(
            f"{item.get('brand', '')} {item.get('title', '')}"
        ):
            neg_counts[token] = neg_counts.get(token, 0) + 1

    # Net weights: approvals boost, skips dampen
    weights = {}
    for token, count in pos_counts.items():
        weights[token] = count - 0.5 * neg_counts.get(token, 0)
    for token, count in neg_counts.items():
        if token not in weights:
            weights[token] = -0.5 * count

    ranked = sorted(weights.items(), key=lambda x: x[1], reverse=True)
    positive = [t for t, w in ranked if w > 0][:25]
    negative = [t for t, w in ranked if w < 0][:15]
    top_brands = sorted(brand_counts.items(), key=lambda x: x[1], reverse=True)

    # Learned search queries: brand + tokens that co-occurred with that brand
    brand_token_map: dict[str, dict[str, int]] = {}
    for item in approved:
        brand = item.get("brand", "unknown")
        brand_token_map.setdefault(brand, {})
        for token in extract_taste_tokens(f"{brand} {item.get('title', '')}"):
            brand_token_map[brand][token] = brand_token_map[brand].get(token, 0) + 1

    learned_queries = []
    for brand, _ in top_brands[:3]:
        brand_seed = (
            brand.lower()
            .replace("galliano dior", "dior galliano")
            .replace("tom ford gucci", "gucci tom ford")
        )
        token_rank = sorted(
            brand_token_map.get(brand, {}).items(),
            key=lambda x: x[1],
            reverse=True,
        )
        for token, _ in token_rank[:6]:
            if token in brand_seed:
                continue
            parts = []
            for part in f"{brand_seed} {token}".split():
                if part not in parts:
                    parts.append(part)
            q = " ".join(parts).strip()
            if len(q.split()) < 2:
                continue
            if q not in learned_queries:
                learned_queries.append(q)
            if len(learned_queries) >= 10:
                break

    avg_price = round(sum(prices) / len(prices), 0) if prices else None

    # Seed / boost with explicit aesthetic north-star tokens
    inspiration_boosts = {
        "chanel": 3, "2.55": 3, "classic flap": 2, "rare": 3, "archive": 2,
        "saddle": 3, "trotter": 3, "dior": 2, "heels": 2, "bag": 2,
        "galliano": 3, "cavalli": 3, "leopard": 3, "animal print": 3,
        "gaultier": 3, "tom ford": 3, "gucci": 2, "jackie": 2, "y2k": 2,
    }
    for token, boost in inspiration_boosts.items():
        weights[token] = weights.get(token, 0) + boost
    ranked = sorted(weights.items(), key=lambda x: x[1], reverse=True)
    positive = [t for t, w in ranked if w > 0][:25]
    negative = [t for t, w in ranked if w < 0][:15]

    global _TASTE_CACHE
    profile = {
        "updated_at": datetime.now().isoformat(),
        "approved_count": len(approved),
        "skipped_count": len(skipped),
        "positive_tokens": positive,
        "negative_tokens": negative,
        "token_weights": {k: round(v, 2) for k, v in ranked[:40]},
        "brand_counts": brand_counts,
        "avg_approved_price": avg_price,
        "learned_queries": learned_queries,
        "inspiration": [
            "Isle of Monday",
            "Jean Vintage",
            "Break Archive",
            "Chanel rare bags",
            "Dior rare bags / shoes / clothes",
            "Cavalli • Galliano • Tom Ford Gucci • Gaultier",
        ],
        "ready": len(approved) >= TASTE_MIN_SAMPLES,
    }
    save_json(TASTE_DB, profile)
    _TASTE_CACHE = profile
    log.info(
        "Taste profile retrained on %s approvals / %s skips. ready=%s",
        len(approved), len(skipped), profile["ready"],
    )
    return profile


def load_taste_profile() -> dict:
    global _TASTE_CACHE
    if _TASTE_CACHE is not None:
        return _TASTE_CACHE
    profile = load_json(TASTE_DB, {})
    if not profile:
        return rebuild_taste_profile()
    _TASTE_CACHE = profile
    return profile


def taste_score(title: str, brand: str = "", extra: str = "") -> float:
    """Score how close a listing is to the learned aesthetic."""
    profile = load_taste_profile()
    if not profile.get("ready"):
        return 0.0
    hay = expand_cjk_hay(f"{brand} {title} {extra}".lower())
    weights = profile.get("token_weights", {})
    score = 0.0
    for token, weight in weights.items():
        if token in hay:
            score += float(weight)
    return score


def matches_rule(title: str, price: float, condition: str, rule: dict, extra: str = "", platform: str | None = None) -> bool:
    haystack = expand_cjk_hay(f"{title} {extra}".lower())
    # keywords may include Japanese / Chinese; latin aliases added by expand_cjk_hay
    if not any(kw.lower() in haystack for kw in rule["keywords"]):
        return False
    min_price, max_price = effective_price_bounds(rule, platform=platform)
    if not (min_price <= price <= max_price):
        return False
    if not condition_matches_rule(condition, rule):
        return False

    cleaned = (
        haystack
        .replace("womenswear", "women")
        .replace("women's", "women")
        .replace("womens", "women")
        .replace("women-", "women ")
    )
    exclude = [x.lower() for x in rule.get("exclude", [])]
    if exclude and any(term in cleaned for term in exclude):
        return False

    include = [x.lower() for x in rule.get("include", [])]
    if include and not any(term in cleaned for term in include):
        return False

    style_terms = [x.lower() for x in rule.get("style", [])]
    era_terms = [x.lower() for x in rule.get("era", [])]
    has_style = any(term in cleaned for term in style_terms)
    has_era = any(term in cleaned for term in era_terms)
    hot_shape = any(term in cleaned for term in [
        "mini", "bodycon", "halter", "corset", "bustier", "cutout",
        "slip dress", "lace", "mesh", "sheer", "ruched",
        "bag", "clutch", "heel", "pump", "sandal", "mule", "boot",
        "top", "blouse", "camisole", "knit", "deconstructed",
        *CJK_SHAPE_INCLUDE,
    ])

    women_signal = "women" in cleaned or "femme" in cleaned
    jp_shape = any(term in haystack for term in CJK_SHAPE_INCLUDE)
    if not ((has_style or has_era) and (hot_shape or women_signal or jp_shape)):
        return False

    # Category focus: bags / shoes / clothes stay in their lane
    focus = rule.get("focus")
    if focus in {"bags", "bags_shoes"}:
        bag_shoe = any(term in cleaned for term in [
            "bag", "handbag", "clutch", "tote", "purse", "saddle", "trotter",
            "flap", "2.55", "jackie",
            "heel", "pump", "sandal", "mule", "stiletto", "boot", "shoe",
            "バッグ", "ヒール", "パンプス", "ミュール", "サンダル", "靴",
            "包", "手袋", "高跟鞋", "跟鞋", "穆勒鞋", "涼鞋", "凉鞋", "鞋",
        ])
        if focus == "bags":
            bag_hit = any(term in cleaned for term in BAGS_INCLUDE + [
                "バッグ", "ハンドバッグ", "クラッチ", "包", "手袋", "手拿包",
                "馬鞍包", "马鞍包",
            ])
            if not bag_hit:
                return False
        elif focus == "bags_shoes" and not bag_shoe:
            return False
    elif focus == "shoes":
        shoe_hit = any(term in cleaned for term in SHOES_INCLUDE + [
            "ヒール", "パンプス", "ミュール", "サンダル", "靴",
            "高跟鞋", "跟鞋", "穆勒鞋", "涼鞋", "凉鞋", "鞋",
        ])
        if not shoe_hit:
            return False
    elif focus == "clothes":
        clothes_hit = any(term in cleaned for term in CLOTHES_INCLUDE + CJK_SHAPE_INCLUDE)
        # Tom Ford Gucci may also surface Jackie / bamboo bags
        if rule.get("brand") == "Tom Ford Gucci":
            ford_era = any(term in cleaned for term in [
                "tom ford", "by tom ford", "jackie", "bamboo", "horsebit",
                "1995", "1996", "1997", "1998", "1999", "2000", "2001",
                "2002", "2003", "2004", "トムフォード",
            ])
            bag_hit = any(term in cleaned for term in BAGS_INCLUDE)
            if not ford_era:
                return False
            if not (clothes_hit or bag_hit):
                return False
        elif not clothes_hit:
            return False

    # Once enough Approves exist, require some overlap with learned taste
    profile = load_taste_profile()
    if profile.get("ready"):
        positives = profile.get("positive_tokens", [])[:12]
        negatives = set(profile.get("negative_tokens", [])[:10])
        if any(neg in cleaned for neg in negatives) and not any(
            pos in cleaned for pos in positives[:6]
        ):
            return False
        # Soft gate: prefer items that share at least one trained signal
        if positives and not any(pos in cleaned for pos in positives):
            # Still allow strong archive style hits so discovery doesn't die
            hard_archive = any(term in cleaned for term in [
                "galliano", "columbus", "bondage", "cavalli", "leopard",
                "python", "animal", "saddle", "trotter", "chanel", "2.55",
                "gaultier", "tom ford", "jackie", "bamboo", "horsebit",
                "harness", "rare", "archive",
                "ガリアーノ", "ディオール", "シャネル", "カヴァリ", "ゴルチエ",
                "加利亞諾", "加利亚诺", "迪奧", "迪奥", "香奈兒", "香奈尔",
                "卡沃利", "高緹耶", "高缇耶",
            ])
            if not hard_archive:
                return False

    return True


def url_matches_rule(url: str, rule: dict) -> bool:
    """Cheap prefilter for sitemap URLs before fetching product pages."""
    lowered = url.lower()
    brand_tokens = {
        token.lower().replace(" ", "-")
        for token in [rule["brand"], *rule["keywords"]]
    }
    if not any(token in lowered for token in brand_tokens):
        return False

    exclude = [x.lower().replace(" ", "-") for x in rule.get("exclude", [])]
    if exclude and any(term in lowered for term in exclude):
        return False

    include = [x.lower().replace(" ", "-") for x in rule.get("include", [])]
    feminine_url_hints = [
        "dress", "gown", "skirt", "blouse", "bag", "handbag", "clutch", "tote",
        "purse", "heel", "sandal", "women", "womens", "femme",
    ]
    hints = include + feminine_url_hints
    return any(term in lowered for term in hints)


def condition_matches_rule(condition: str, rule: dict) -> bool:
    """Treat unknown marketplace condition as a soft match."""
    if not condition or condition.lower() == "unknown":
        return True
    return any(c.lower() in condition.lower() for c in rule["conditions"])


def abs_media_url(url: str, base: str = "https://auctions.yahoo.co.jp") -> str:
    url = (url or "").strip()
    if not url or url.startswith("data:") or "placeholder" in url.lower():
        return ""
    if url.startswith("//"):
        return "https:" + url
    if url.startswith("/"):
        return urljoin(base, url)
    if url.startswith("http://"):
        return "https://" + url[len("http://"):]
    return url if url.startswith("http") else ""


_DEAD_LISTING = (
    "this listing has been sold",
    "listing not found",
    "page not found",
    "isn't available",
    "no longer available",
    "終了しました",
    "このオークションは終了",
    "商品は終了",
    "已結束",
    "商品已結束",
    "找不到商品",
)


async def listing_is_live(session: AsyncSession, item: dict) -> bool:
    """Drop ended auctions / sold pages so Telegram doesn't get empty Buyee/Yahoo links."""
    platform = item.get("platform") or ""
    if platform not in {"YahooJP", "Buyee", "YahooTW"}:
        return True
    url = item.get("source_url") or item.get("url") or ""
    if not url:
        return False
    try:
        resp = await session.get(url, headers=HEADERS, timeout=12)
    except Exception:
        return False
    if resp.status_code in (404, 410, 410):
        return False
    if resp.status_code >= 400:
        return False
    text = resp.text or ""
    low = text.lower()
    for marker in _DEAD_LISTING:
        if marker.isascii():
            if marker in low:
                return False
        elif marker in text:
            return False
    return True


async def fetch_photo_bytes(session: AsyncSession, url: str) -> bytes | None:
    """Pull the image ourselves — Telegram's servers are blocked by Grailed/Yahoo CDNs."""
    url = abs_media_url(url)
    if not url:
        return None
    try:
        resp = await session.get(
            url,
            headers={**HEADERS, "Accept": "image/avif,image/webp,image/*,*/*;q=0.8"},
            timeout=15,
        )
        if resp.status_code != 200:
            return None
        data = resp.content or b""
        if len(data) < 2500:
            return None
        return data
    except Exception:
        return None


# ─── Scrapers ─────────────────────────────────────────────────────────────────

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}


async def scrape_grailed(session: AsyncSession, rule: dict) -> list[dict]:
    """Search Grailed via Algolia, split by US / Asia / Europe for price lanes."""
    results = []
    url = (
        f"https://{GRAILED_ALGOLIA_APP_ID}-dsn.algolia.net/1/indexes/"
        f"{GRAILED_LISTING_INDEX}/query"
    )
    headers = {
        "X-Algolia-Application-Id": GRAILED_ALGOLIA_APP_ID,
        "X-Algolia-API-Key": GRAILED_ALGOLIA_SEARCH_KEY,
        "Content-Type": "application/json",
    }
    queries = rule.get("search_queries") or [
        f'{rule["keywords"][0]} dress',
        f'{rule["keywords"][0]} bag clutch',
    ]
    profile = load_taste_profile()
    if profile.get("ready"):
        for q in profile.get("learned_queries", [])[:5]:
            if q not in queries:
                queries.append(q)
    cjk_queries = cjk_search_queries(rule)

    regions = rule.get("grailed_regions") or list(GRAILED_REGIONS)

    try:
        listings = []
        for query in queries:
            for loc_name, loc_code in regions:
                payload = {
                    "query": query,
                    "hitsPerPage": 28,
                    "page": 0,
                    "facetFilters": [[f"location:{loc_name}"]],
                }
                resp = await session.post(
                    url,
                    headers=headers,
                    data=json.dumps(payload),
                    timeout=HTTP_TIMEOUT_SECONDS,
                )
                if resp.status_code == 200:
                    for hit in resp.json().get("hits", []):
                        hit = dict(hit)
                        hit["_region"] = loc_code
                        hit["_region_label"] = loc_name
                        listings.append(hit)
                await asyncio.sleep(0.2)
        # Native-script queries: Asia sellers often title in JP/CN
        for query in cjk_queries[:8]:
            payload = {
                "query": query,
                "hitsPerPage": 28,
                "page": 0,
                "facetFilters": [["location:Asia"]],
            }
            resp = await session.post(
                url,
                headers=headers,
                data=json.dumps(payload),
                timeout=HTTP_TIMEOUT_SECONDS,
            )
            if resp.status_code == 200:
                for hit in resp.json().get("hits", []):
                    hit = dict(hit)
                    hit["_region"] = "Asia"
                    hit["_region_label"] = "Asia"
                    listings.append(hit)
            await asyncio.sleep(0.2)

        seen_ids = set()
        for item in listings[:200]:
            listing_id = str(item.get("id", ""))
            if listing_id in seen_ids:
                continue
            seen_ids.add(listing_id)

            title = item.get("title", "")
            price = float(item.get("price_i", 0))
            condition = item.get("condition", "")
            if item.get("sold") or item.get("sold_at") or str(item.get("status") or "").lower() == "sold":
                continue
            photo = abs_media_url(
                item.get("cover_photo", {}).get("url")
                or item.get("cover_photo", {}).get("image_url", "")
            )
            slug = item.get("slug", listing_id)
            region = item.get("_region") or "—"
            region_label = item.get("_region_label") or item.get("location") or region
            extra = " ".join(
                filter(
                    None,
                    [
                        item.get("department", ""),
                        item.get("category", ""),
                        item.get("category_path", ""),
                        item.get("designer_names", ""),
                        str(region_label),
                    ],
                )
            )

            if matches_rule(title, price, condition, rule, extra=extra):
                results.append({
                    "platform": "Grailed",
                    "id": listing_id,
                    "brand": rule["brand"],
                    "title": title,
                    "price": price,
                    "condition": condition,
                    "photo": photo,
                    "heat": float(item.get("heat_f") or item.get("heat") or 0),
                    "taste": taste_score(title, rule["brand"], extra),
                    "url": (
                        f"https://www.grailed.com/listings/{slug}"
                        if slug
                        else item.get("url", "")
                    ),
                    "region": region,
                    "region_label": region_label,
                })
    except Exception as e:
        log.warning(f"Grailed scrape error: {e}")
    return results


async def fetch_market_heat(session: AsyncSession, query: str) -> dict:
    """Sample Grailed heat for a trend query."""
    url = (
        f"https://{GRAILED_ALGOLIA_APP_ID}-dsn.algolia.net/1/indexes/"
        f"{GRAILED_HEAT_INDEX}/query"
    )
    headers = {
        "X-Algolia-Application-Id": GRAILED_ALGOLIA_APP_ID,
        "X-Algolia-API-Key": GRAILED_ALGOLIA_SEARCH_KEY,
        "Content-Type": "application/json",
    }
    payload = {"query": query, "hitsPerPage": 20, "page": 0}
    try:
        resp = await session.post(
            url, headers=headers, data=json.dumps(payload), timeout=HTTP_TIMEOUT_SECONDS
        )
        if resp.status_code != 200:
            return {"query": query, "count": 0, "avg_heat": 0, "avg_price": 0}
        hits = resp.json().get("hits", [])
        heats = [float(h.get("heat_f") or 0) for h in hits]
        prices = [float(h.get("price_i") or 0) for h in hits if h.get("price_i")]
        return {
            "query": query,
            "count": len(hits),
            "avg_heat": round(sum(heats) / len(heats), 1) if heats else 0,
            "avg_price": round(sum(prices) / len(prices), 0) if prices else 0,
        }
    except Exception as e:
        log.warning("Trend heat fetch failed for %s: %s", query, e)
        return {"query": query, "count": 0, "avg_heat": 0, "avg_price": 0}


def analyze_personal_taste(approved: list) -> list[tuple[str, int]]:
    """Rank trend signals from what you've actually approved."""
    counts = {signal: 0 for signal in TREND_SIGNALS}
    for item in approved:
        hay = f"{item.get('brand', '')} {item.get('title', '')}".lower()
        for signal in TREND_SIGNALS:
            if signal in hay:
                counts[signal] += 1
    ranked = sorted(
        ((k, v) for k, v in counts.items() if v > 0),
        key=lambda x: x[1],
        reverse=True,
    )
    return ranked[:8]


async def build_trend_report() -> str:
    """Combine personal approvals with live Grailed demand heat."""
    approved = load_json(APPROVED_DB, [])
    personal = analyze_personal_taste(approved)

    trend_queries = [
        "galliano dior",
        "dior bondage",
        "dior columbus",
        "tom ford gucci",
        "gucci jackie tom ford",
        "cavalli leopard dress",
        "y2k archive dress",
    ]

    async with AsyncSession(impersonate="chrome", headers=HEADERS) as session:
        market = []
        for q in trend_queries:
            market.append(await fetch_market_heat(session, q))
            await asyncio.sleep(0.25)

    market_ranked = sorted(market, key=lambda x: x["avg_heat"], reverse=True)
    save_json(TREND_DB, {
        "updated_at": datetime.now().isoformat(),
        "personal": personal,
        "market": market_ranked,
    })

    lines = ["📈 Fashion demand snapshot\n"]
    settings = load_settings()
    gmin = settings.get("global_min_price")
    gmax = settings.get("global_max_price")
    if gmin is not None or gmax is not None:
        lines.append(
            f"Active price filter: ${gmin or 0:.0f}–${gmax or '∞'}\n"
        )

    lines.append("Your taste (from /approved):")
    if personal:
        for signal, count in personal:
            lines.append(f"• {signal}: {count} saves")
    else:
        lines.append("• Approve a few pieces first to train this.")

    lines.append("\nMarket heat on Grailed (higher = more demand right now):")
    for row in market_ranked[:6]:
        lines.append(
            f"• {row['query']}: heat {row['avg_heat']:.0f} • avg ${row['avg_price']:.0f}"
        )

    # Simple recommendation: intersection of personal + hot market
    hot_words = " ".join(r["query"] for r in market_ranked[:3]).lower()
    tips = []
    for signal, _ in personal[:5]:
        if signal in hot_words or any(signal in r["query"] for r in market_ranked[:4]):
            tips.append(signal)
    if not tips:
        tips = [r["query"] for r in market_ranked[:3]]

    lines.append("\nBuy focus now:")
    for tip in tips[:4]:
        lines.append(f"→ {tip}")

    return "\n".join(lines)


async def fetch_vestiaire_item(session: AsyncSession, url: str) -> dict | None:
    """Fetch a Vestiaire product page and extract schema.org product data."""
    try:
        resp = await session.get(url, headers=HEADERS, timeout=HTTP_TIMEOUT_SECONDS)
        if resp.status_code != 200:
            return None
    except Exception:
        return None

    soup = BeautifulSoup(resp.text, "html.parser")
    for script in soup.find_all("script", type="application/ld+json"):
        raw = (script.string or script.get_text() or "").strip()
        if not raw:
            continue
        try:
            payload = json.loads(raw)
        except json.JSONDecodeError:
            continue

        entries = payload if isinstance(payload, list) else [payload]
        for entry in entries:
            if not isinstance(entry, dict) or entry.get("@type") != "Product":
                continue

            offers = entry.get("offers", {}) or {}
            brand = entry.get("brand", {}) or {}
            image = entry.get("image", {}) or {}

            try:
                price = float(offers.get("price", 0))
            except (TypeError, ValueError):
                continue

            condition_url = offers.get("itemCondition", "")
            condition_token = condition_url.rsplit("/", 1)[-1].replace("Condition", "")
            condition = condition_token or "Unknown"
            if condition == "New":
                condition = "Excellent"
            elif condition == "Used":
                condition = "Unknown"
            photo = image.get("contentUrl") if isinstance(image, dict) else ""

            return {
                "platform": "Vestiaire",
                "id": str(entry.get("sku", make_id(url))),
                "brand_name": brand.get("name", ""),
                "title": entry.get("name", ""),
                "price": price,
                "condition": condition,
                "photo": photo,
                "url": offers.get("url", url) or url,
            }
    return None


async def scrape_vestiaire(session: AsyncSession, rule: dict) -> list[dict]:
    """Discover Vestiaire products from public sitemaps, then filter locally."""
    results = []
    try:
        resp = await session.get(VESTIAIRE_SITEMAP_INDEX, headers=HEADERS, timeout=HTTP_TIMEOUT_SECONDS)
        if resp.status_code != 200:
            return results
        sitemap_index = BeautifulSoup(resp.text, "xml")
        sitemap_urls = [
            loc.get_text(strip=True)
            for loc in sitemap_index.find_all("loc")
            if "products-chunk" in loc.get_text(strip=True)
        ][:VESTIAIRE_PRODUCT_SITEMAP_LIMIT]

        sitemap_sem = asyncio.Semaphore(VESTIAIRE_SITEMAP_FETCH_CONCURRENCY)

        async def fetch_matching_urls(sitemap_url: str) -> list[str]:
            async with sitemap_sem:
                sitemap_resp = await session.get(
                    sitemap_url,
                    headers=HEADERS,
                    timeout=HTTP_TIMEOUT_SECONDS,
                )
            if sitemap_resp.status_code != 200:
                return []
            sitemap = BeautifulSoup(sitemap_resp.text, "xml")
            return [
                loc.get_text(strip=True)
                for loc in sitemap.find_all("loc")
                if loc.get_text(strip=True).endswith(".shtml")
                and url_matches_rule(loc.get_text(strip=True), rule)
            ]

        product_urls = []
        sitemap_matches = await asyncio.gather(
            *(fetch_matching_urls(sitemap_url) for sitemap_url in sitemap_urls),
            return_exceptions=True,
        )
        for match_group in sitemap_matches:
            if isinstance(match_group, Exception):
                continue
            product_urls.extend(match_group)
            if len(product_urls) >= VESTIAIRE_MATCHING_URL_LIMIT:
                break

        seen_urls = set()
        deduped_urls = []
        for url in product_urls:
            if url not in seen_urls:
                seen_urls.add(url)
                deduped_urls.append(url)
            if len(deduped_urls) >= VESTIAIRE_MATCHING_URL_LIMIT:
                break

        if not deduped_urls:
            return results

        sem = asyncio.Semaphore(VESTIAIRE_PRODUCT_FETCH_CONCURRENCY)

        async def fetch_one(product_url: str):
            async with sem:
                return await fetch_vestiaire_item(session, product_url)

        items = await asyncio.gather(*(fetch_one(url) for url in deduped_urls), return_exceptions=True)

        for item in items:
            if isinstance(item, Exception) or not item:
                continue

            haystack = f"{item['brand_name']} {item['title']}".strip()
            if matches_rule(haystack, item["price"], item["condition"], rule, extra=item.get("url", "")):
                title = f"{item['brand_name']} {item['title']}".strip()
                results.append({
                    "platform":  "Vestiaire",
                    "id":        item["id"],
                    "brand":     rule["brand"],
                    "title":     title,
                    "price":     item["price"],
                    "condition": item["condition"],
                    "photo":     item["photo"],
                    "taste":     taste_score(title, rule["brand"], item.get("url", "")),
                    "url":       item["url"],
                    "region": "EU",
                    "region_label": "Europe / Italy lane",
                })
    except Exception as e:
        log.warning(f"Vestiaire scrape error: {e}")
    return results


async def scrape_yahoo_jp(session: AsyncSession, rule: dict) -> list[dict]:
    """
    Search Yahoo Auctions Japan directly (cheaper archive lane).
    Purchase links go through Buyee so you can still buy from abroad.
    """
    results = []
    seen_auction_ids = set()
    queries = [q for q in (rule.get("jp_queries") or []) if q]
    min_usd, max_usd = effective_price_bounds(rule, platform="YahooJP")
    min_jpy = int(min_usd * JPY_PER_USD)
    max_jpy = int(max_usd * JPY_PER_USD)

    for query in queries:
        if not query:
            continue
        url = (
            "https://auctions.yahoo.co.jp/search/search"
            f"?p={quote(query)}&va={quote(query)}&exflg=1&b=1&n=50"
            f"&aucminprice={min_jpy}&aucmaxprice={max_jpy}"
            "&price_type=currentprice&istatus=1"
        )
        try:
            resp = await session.get(url, headers=HEADERS, timeout=HTTP_TIMEOUT_SECONDS)
            if resp.status_code != 200:
                log.warning("YahooJP status %s for %s", resp.status_code, query)
                continue
            soup = BeautifulSoup(resp.text, "html.parser")
            products = soup.select("li.Product")
            if not products:
                continue

            for product in products[:30]:
                blob = product.get_text(" ", strip=True)
                if "終了" in blob:
                    continue
                title_el = product.select_one(".Product__titleLink")
                price_el = product.select_one(".Product__priceValue") or product.select_one(".Product__price")
                photo_el = product.select_one("img")
                if not (title_el and price_el):
                    continue

                title = title_el.get_text(strip=True)
                href = (title_el.get("href") or "").split("?")[0]
                if href and not href.startswith("http"):
                    href = urljoin("https://page.auctions.yahoo.co.jp/", href)
                auction_id = href.rstrip("/").split("/")[-1] if href else ""
                if not auction_id or auction_id in seen_auction_ids:
                    continue
                seen_auction_ids.add(auction_id)
                photo = ""
                if photo_el:
                    photo = abs_media_url(
                        photo_el.get("src")
                        or photo_el.get("data-src")
                        or photo_el.get("data-original")
                        or ""
                    )

                try:
                    price_jpy = float(
                        price_el.get_text(strip=True)
                        .replace(",", "")
                        .replace("円", "")
                        .replace("¥", "")
                    )
                    price_usd = price_jpy / JPY_PER_USD
                except Exception:
                    continue

                # Buyee proxy checkout for overseas buyers
                buyee_url = f"https://buyee.jp/item/jdirectitems/auction/{auction_id}"

                if matches_rule(title, price_usd, "B", rule, platform="YahooJP"):
                    results.append({
                        "platform":  "YahooJP",
                        "id":        auction_id,
                        "brand":     rule["brand"],
                        "title":     title,
                        "price":     round(price_usd, 2),
                        "condition": "JP auction",
                        "photo":     photo,
                        "taste":     taste_score(title, rule["brand"]),
                        "url":       buyee_url,
                        "source_url": href,
                        "region": "JP",
                        "region_label": "Japan",
                    })
            await asyncio.sleep(0.4)
        except Exception as e:
            log.warning("YahooJP scrape error (%s): %s", query, e)

    return results


async def scrape_yahoo_tw(session: AsyncSession, rule: dict) -> list[dict]:
    """Search Yahoo Auctions Taiwan — often priced below US / EU comps."""
    results = []
    seen_ids: set[str] = set()
    queries = list(dict.fromkeys([
        *(rule.get("tw_queries") or []),
        *(rule.get("cn_queries") or []),
        *(rule.get("search_queries") or [])[:1],
    ]))
    min_usd, max_usd = effective_price_bounds(rule, platform="YahooTW")
    min_twd = int(min_usd * TWD_PER_USD)
    max_twd = int(max_usd * TWD_PER_USD)

    for query in queries:
        if not query:
            continue
        url = f"https://tw.bid.yahoo.com/search/auction/product?p={quote(query)}"
        try:
            resp = await session.get(url, headers=HEADERS, timeout=HTTP_TIMEOUT_SECONDS)
            if resp.status_code != 200:
                log.warning("YahooTW status %s for %s", resp.status_code, query)
                continue
            soup = BeautifulSoup(resp.text, "html.parser")
            for link in soup.select('a[href*="/item/"]')[:40]:
                href = link.get("href") or ""
                if "/item/" not in href:
                    continue
                item_id = href.rstrip("/").split("/")[-1].split("?")[0]
                if not item_id or item_id in seen_ids:
                    continue
                title = link.get_text(" ", strip=True)
                if not title or len(title) < 6:
                    continue
                # Prefer price from the link text: "$ 54,000" is TWD on Yahoo TW
                money = re.findall(r"\$\s*([\d,]+)", title)
                if not money:
                    parent = link.find_parent(["li", "div", "article"])
                    blob = parent.get_text(" ", strip=True) if parent else title
                    money = re.findall(r"\$\s*([\d,]+)", blob)
                if not money:
                    continue
                try:
                    # Use the lower of first two prices when a strike + sale show up
                    prices_twd = [float(m.replace(",", "")) for m in money[:2]]
                    price_twd = min(prices_twd)
                    if price_twd < min_twd or price_twd > max_twd:
                        # Still let matches_rule decide with USD conversion
                        pass
                    price_usd = price_twd / TWD_PER_USD
                except Exception:
                    continue

                seen_ids.add(item_id)
                # Drop noise words from title for matching (prices etc.)
                clean_title = re.sub(r"\$\s*[\d,]+", " ", title)
                clean_title = re.sub(r"\s+", " ", clean_title).strip()

                if "結束" in clean_title or "已售" in clean_title:
                    continue
                parent = link.find_parent(["li", "div", "article"])
                img = (parent.select_one("img") if parent else None) or link.select_one("img")
                photo = abs_media_url(
                    (img.get("src") or img.get("data-src") or "") if img else ""
                )
                listing_url = href if href.startswith("http") else f"https://tw.bid.yahoo.com/item/{item_id}"
                if matches_rule(clean_title, price_usd, "B", rule, platform="YahooTW"):
                    results.append({
                        "platform": "YahooTW",
                        "id": item_id,
                        "brand": rule["brand"],
                        "title": clean_title[:180],
                        "price": round(price_usd, 2),
                        "condition": "TW auction",
                        "photo": photo,
                        "taste": taste_score(clean_title, rule["brand"]),
                        "url": listing_url,
                        "source_url": listing_url,
                        "region": "TW",
                        "region_label": "Taiwan",
                    })
            await asyncio.sleep(0.45)
        except Exception as e:
            log.warning("YahooTW scrape error (%s): %s", query, e)

    return results


async def scrape_buyee(session: AsyncSession, rule: dict) -> list[dict]:
    """Legacy Buyee HTML search — usually WAF-blocked; YahooJP is the primary Japan lane."""
    return await scrape_yahoo_jp(session, rule)

# ─── Core Agent Loop ──────────────────────────────────────────────────────────

async def run_sourcing_scan(app):
    """Main scan loop — called by scheduler every POLL_INTERVAL minutes."""
    seen    = load_json(SEEN_DB, {})
    pending = load_json(PENDING_DB, {})
    found   = 0

    async with AsyncSession(impersonate="chrome", headers=HEADERS) as session:
        for rule in RULES:
            log.info(f"Scanning for {rule['brand']}...")
            # Run platforms one-by-one so Telegram polling stays responsive
            all_results = []
            for scraper in (
                scrape_grailed,
                scrape_vestiaire,
                scrape_yahoo_jp,
                scrape_yahoo_tw,
            ):
                try:
                    batch = await scraper(session, rule)
                    all_results.extend(batch)
                except Exception as e:
                    log.warning("Scraper task failed: %s", e)
                await asyncio.sleep(0.2)

            all_results.sort(
                key=lambda x: (
                    float(x.get("taste") or 0),
                    float(x.get("heat") or 0),
                ),
                reverse=True,
            )

            for item in all_results:
                uid = f"{item['platform']}_{item['id']}"
                if uid in seen:
                    continue

                live = await listing_is_live(session, item)
                if not live:
                    seen[uid] = datetime.now().isoformat()
                    save_json(SEEN_DB, seen)
                    log.info("Skip stale listing %s", uid)
                    continue

                seen[uid]    = datetime.now().isoformat()
                pending[uid] = normalize_item_brand(item)
                found += 1

                # Persist immediately so Approve/Skip works during a long scan
                save_json(SEEN_DB, seen)
                save_json(PENDING_DB, pending)

                await send_item_to_telegram(app, item, uid, session)
                await asyncio.sleep(1)  # avoid rate limiting

    save_json(SEEN_DB, seen)
    save_json(PENDING_DB, pending)
    log.info(f"Scan complete. {found} new items found.")


async def send_item_to_telegram(app, item: dict, uid: str, session: AsyncSession | None = None):
    """Send a sourced item to Telegram with Approve / Skip buttons."""
    condition_emoji = {
        "Excellent": "🟢", "Very Good": "🟡", "Good": "🟠",
        "Fair": "🔴", "A": "🟢", "B": "🟡", "C": "🟠",
        "is_new": "✨", "gently_used": "🟡", "is_gently_used": "🟡", "used": "🟠",
        "is_used": "🟠", "Unknown": "⚪",
    }
    cond_icon = next(
        (v for k, v in condition_emoji.items() if k.lower() in item["condition"].lower()),
        "⚪"
    )

    region = item.get("region_label") or item.get("region") or ""
    region_line = f"📍 Region: {region}\n" if region else ""
    listing_url = item.get("url") or ""
    text = (
        "feling. sourcing agent\n\n"
        f"🏷 {item['brand']}\n"
        f"📝 {item['title']}\n\n"
        f"💰 ${item['price']:.0f} USD\n"
        f"{cond_icon} Condition: {item['condition']}\n"
        f"🌐 Platform: {item['platform']}\n"
        f"{region_line}\n"
        f"{listing_url}"
    )

    rows = []
    if listing_url.startswith("http"):
        rows.append([InlineKeyboardButton("👀 Open listing", url=listing_url)])
    rows.append([
        InlineKeyboardButton("✅ Approve — buy this", callback_data=f"approve:{uid}"),
        InlineKeyboardButton("❌ Skip", callback_data=f"skip:{uid}"),
    ])
    keyboard = InlineKeyboardMarkup(rows)

    photo_bytes = None
    if session is not None and item.get("photo"):
        photo_bytes = await fetch_photo_bytes(session, item["photo"])

    try:
        if photo_bytes:
            await app.bot.send_photo(
                chat_id=TELEGRAM_CHAT_ID,
                photo=InputFile(BytesIO(photo_bytes), filename=f"{uid}.jpg"),
                caption=text[:1024],
                reply_markup=keyboard,
                read_timeout=20,
                write_timeout=20,
            )
            return
        await app.bot.send_message(
            chat_id=TELEGRAM_CHAT_ID,
            text=text,
            reply_markup=keyboard,
            disable_web_page_preview=False,
        )
    except Exception as e:
        log.error(f"Telegram send error: {e}")


async def append_message_status(query, status_text: str):
    """Append a status line to either a photo caption or text message."""
    message = query.message
    if message.caption is not None:
        await query.edit_message_caption(
            caption=(message.caption + status_text)[:1024],
            reply_markup=None,
        )
    elif message.text is not None:
        await query.edit_message_text(
            text=message.text + status_text,
            reply_markup=None,
            disable_web_page_preview=False,
        )

# ─── Weekly approved report (comps → cost table → email) ──────────────────────

def condition_rank(condition: str) -> int:
    key = (condition or "unknown").lower().strip()
    if key in CONDITION_RANK:
        return CONDITION_RANK[key]
    for name, rank in CONDITION_RANK.items():
        if name in key:
            return rank
    return 0


def approved_in_lookback(days: int = WEEKLY_REPORT_LOOKBACK_DAYS) -> list[dict]:
    approved = load_json(APPROVED_DB, [])
    cutoff = datetime.now() - timedelta(days=days)
    recent = []
    for item in approved:
        raw = item.get("approved_at") or ""
        try:
            ts = datetime.fromisoformat(raw)
        except ValueError:
            ts = None
        if ts is None or ts >= cutoff:
            recent.append(item)
    # Prefer recent stamps; fall back to last N if dates sparse
    if not recent:
        recent = approved[-WEEKLY_REPORT_MAX_ITEMS:]
    return recent[-WEEKLY_REPORT_MAX_ITEMS:]


def build_comp_query(item: dict) -> str:
    """Build a short marketplace query from an approved listing."""
    brand = re.sub(r"[^a-zA-Z0-9\s]", " ", (item.get("brand") or "")).lower()
    brand = re.sub(
        r"\b(rare bags|bags|shoes|clothes|tom ford gucci|bags & shoes)\b",
        " ",
        brand,
    )
    title = (item.get("title") or "").lower()
    # Keep brand-ish + concrete product words
    keep = []
    for word in re.findall(r"[a-z0-9']{3,}", f"{brand} {title}"):
        if word in STOP_WORDS:
            continue
        if word.isdigit():
            continue
        if word not in keep:
            keep.append(word)
        if len(keep) >= 7:
            break
    return " ".join(keep) or (item.get("brand") or "vintage")


async def search_grailed_comps(session: AsyncSession, query: str, limit: int = 20) -> list[dict]:
    """Open search on Grailed Algolia for comparable listings (no rule filter)."""
    url = (
        f"https://{GRAILED_ALGOLIA_APP_ID}-dsn.algolia.net/1/indexes/"
        f"{GRAILED_LISTING_INDEX}/query"
    )
    headers = {
        "X-Algolia-Application-Id": GRAILED_ALGOLIA_APP_ID,
        "X-Algolia-API-Key": GRAILED_ALGOLIA_SEARCH_KEY,
        "Content-Type": "application/json",
    }
    results = []
    try:
        resp = await session.post(
            url,
            headers=headers,
            data=json.dumps({"query": query, "hitsPerPage": limit, "page": 0}),
            timeout=HTTP_TIMEOUT_SECONDS,
        )
        if resp.status_code != 200:
            return results
        for hit in resp.json().get("hits", []):
            listing_id = str(hit.get("id", ""))
            slug = hit.get("slug", listing_id)
            results.append({
                "platform": "Grailed",
                "id": listing_id,
                "title": hit.get("title", ""),
                "price": float(hit.get("price_i", 0) or 0),
                "condition": hit.get("condition", "unknown"),
                "url": f"https://www.grailed.com/listings/{slug}" if slug else "",
                "heat": float(hit.get("heat_f") or hit.get("heat") or 0),
            })
    except Exception as e:
        log.warning("Comp search failed for %r: %s", query, e)
    return results


def title_similarity(a: str, b: str) -> float:
    ta = set(re.findall(r"[a-z0-9']{3,}", (a or "").lower())) - STOP_WORDS
    tb = set(re.findall(r"[a-z0-9']{3,}", (b or "").lower())) - STOP_WORDS
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


def pick_better_comp(approved: dict, comps: list[dict]) -> dict | None:
    """
    Before tabulating: find a meaningfully better market option.
    Prefer lower price at similar quality, or better condition at similar price.
    """
    own_id = str(approved.get("id", ""))
    own_url = (approved.get("url") or "").rstrip("/")
    own_price = float(approved.get("price") or 0)
    own_cond = condition_rank(approved.get("condition", ""))
    own_title = approved.get("title", "")

    best = None
    best_score = 0.0
    for comp in comps:
        if str(comp.get("id", "")) == own_id:
            continue
        if (comp.get("url") or "").rstrip("/") == own_url:
            continue
        c_price = float(comp.get("price") or 0)
        if c_price <= 0:
            continue
        sim = title_similarity(own_title, comp.get("title", ""))
        if sim < 0.22:
            continue
        c_cond = condition_rank(comp.get("condition", ""))
        cheaper = own_price > 0 and c_price <= own_price * 0.92
        better_cond = c_cond > own_cond and c_price <= own_price * 1.08
        similar_deal = abs(c_price - own_price) / max(own_price, 1) <= 0.05 and c_cond >= own_cond
        if not (cheaper or better_cond or similar_deal):
            continue
        savings = max(0.0, own_price - c_price)
        score = (
            sim * 4
            + (savings / max(own_price, 1)) * 5
            + max(0, c_cond - own_cond) * 1.5
            + (0.3 if better_cond else 0)
        )
        if score > best_score:
            best_score = score
            best = {
                **comp,
                "similarity": round(sim, 2),
                "savings": round(savings, 0),
                "why": (
                    f"${savings:.0f} cheaper"
                    if cheaper and savings >= 1
                    else (
                        "better condition"
                        if better_cond
                        else "similar price, equal/better condition"
                    )
                ),
            }
    return best


async def scour_approved_comps(items: list[dict]) -> list[dict]:
    """Re-source comps for each approved item before building the cost table."""
    rows = []
    async with AsyncSession(impersonate="chrome", headers=HEADERS) as session:
        for item in items:
            query = build_comp_query(item)
            comps = await search_grailed_comps(session, query, limit=25)
            better = pick_better_comp(item, comps)
            cheaper_alts = [
                c for c in comps
                if c.get("id") != item.get("id")
                and float(c.get("price") or 0) > 0
                and float(c.get("price") or 0) < float(item.get("price") or 0)
                and title_similarity(item.get("title", ""), c.get("title", "")) >= 0.2
            ]
            cheaper_alts.sort(key=lambda x: float(x.get("price") or 0))
            market_low = float(cheaper_alts[0]["price"]) if cheaper_alts else None
            rows.append({
                "item": item,
                "query": query,
                "better": better,
                "market_low": market_low,
                "comp_count": len(comps),
            })
            await asyncio.sleep(0.25)
    return rows


def report_verdict(row: dict) -> str:
    item = row["item"]
    better = row.get("better")
    price = float(item.get("price") or 0)
    market_low = row.get("market_low")
    if better and float(better.get("savings") or 0) >= max(25, price * 0.1):
        return "SWITCH"
    if better and better.get("why") == "better condition":
        return "UPGRADE"
    if market_low is not None and price > 0 and market_low <= price * 0.85:
        return "WATCH"
    return "KEEP"


def build_weekly_report_html(rows: list[dict], window_days: int) -> str:
    total = sum(float(r["item"].get("price") or 0) for r in rows)
    saveable = sum(float((r.get("better") or {}).get("savings") or 0) for r in rows)
    switch_n = sum(1 for r in rows if report_verdict(r) in {"SWITCH", "UPGRADE"})
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")

    body_rows = []
    for r in rows:
        item = r["item"]
        better = r.get("better")
        verdict = report_verdict(r)
        your_price = float(item.get("price") or 0)
        better_cell = "—"
        if better:
            better_cell = (
                f"<a href=\"{html_lib.escape(better.get('url') or '#')}\">"
                f"${float(better.get('price') or 0):.0f}</a>"
                f"<br><span style=\"color:#666;font-size:12px\">"
                f"{html_lib.escape(better.get('why', ''))} · "
                f"{html_lib.escape(str(better.get('condition', '')))}"
                f"</span>"
            )
        market = f"${r['market_low']:.0f}" if r.get("market_low") is not None else "—"
        color = {
            "KEEP": "#1a7f37",
            "WATCH": "#9a6700",
            "UPGRADE": "#0969da",
            "SWITCH": "#cf222e",
        }.get(verdict, "#333")
        body_rows.append(
            "<tr>"
            f"<td style=\"padding:10px;border-bottom:1px solid #eee;max-width:280px\">"
            f"<strong>{html_lib.escape(item.get('brand', ''))}</strong><br>"
            f"<span style=\"font-size:13px\">{html_lib.escape((item.get('title') or '')[:90])}</span><br>"
            f"<a href=\"{html_lib.escape(item.get('url') or '#')}\" style=\"font-size:12px\">open yours</a>"
            f"</td>"
            f"<td style=\"padding:10px;border-bottom:1px solid #eee;text-align:right;font-size:18px;font-weight:700\">"
            f"${your_price:.0f}</td>"
            f"<td style=\"padding:10px;border-bottom:1px solid #eee;text-align:right\">{market}</td>"
            f"<td style=\"padding:10px;border-bottom:1px solid #eee\">{better_cell}</td>"
            f"<td style=\"padding:10px;border-bottom:1px solid #eee;font-weight:700;color:{color}\">{verdict}</td>"
            "</tr>"
        )

    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>feling. weekly approved report</title></head>
<body style="font-family:Georgia,serif;background:#f7f4ef;margin:0;padding:24px;color:#1a1a1a">
  <div style="max-width:920px;margin:0 auto;background:#fff;padding:28px;border:1px solid #e8e0d5">
    <h1 style="margin:0 0 6px;font-size:28px">feling. weekly approved</h1>
    <p style="margin:0 0 20px;color:#666">Last {window_days} days · generated {generated}</p>
    <div style="display:flex;gap:16px;flex-wrap:wrap;margin-bottom:24px">
      <div style="background:#f7f4ef;padding:14px 18px;min-width:140px">
        <div style="font-size:12px;color:#666;text-transform:uppercase">Approved spend</div>
        <div style="font-size:26px;font-weight:700">${total:.0f}</div>
      </div>
      <div style="background:#f7f4ef;padding:14px 18px;min-width:140px">
        <div style="font-size:12px;color:#666;text-transform:uppercase">Potential save</div>
        <div style="font-size:26px;font-weight:700">${saveable:.0f}</div>
      </div>
      <div style="background:#f7f4ef;padding:14px 18px;min-width:140px">
        <div style="font-size:12px;color:#666;text-transform:uppercase">Items</div>
        <div style="font-size:26px;font-weight:700">{len(rows)}</div>
      </div>
      <div style="background:#f7f4ef;padding:14px 18px;min-width:140px">
        <div style="font-size:12px;color:#666;text-transform:uppercase">Switch / upgrade</div>
        <div style="font-size:26px;font-weight:700">{switch_n}</div>
      </div>
    </div>
    <p style="font-size:14px;color:#444;margin:0 0 16px">
      Each row was re-sourced on Grailed for similar pieces <em>before</em> this table.
      Cost is the first column after the item — use SWITCH when a clearer cheaper match exists.
    </p>
    <table style="width:100%;border-collapse:collapse;font-size:14px">
      <thead>
        <tr style="background:#1a1a1a;color:#fff;text-align:left">
          <th style="padding:10px">Item</th>
          <th style="padding:10px;text-align:right">Your cost</th>
          <th style="padding:10px;text-align:right">Market low</th>
          <th style="padding:10px">Better find</th>
          <th style="padding:10px">Call</th>
        </tr>
      </thead>
      <tbody>
        {''.join(body_rows) or '<tr><td colspan="5" style="padding:16px">No approvals in this window.</td></tr>'}
      </tbody>
    </table>
  </div>
</body></html>"""


def build_weekly_report_text(rows: list[dict], window_days: int) -> str:
    total = sum(float(r["item"].get("price") or 0) for r in rows)
    saveable = sum(float((r.get("better") or {}).get("savings") or 0) for r in rows)
    lines = [
        f"feling. weekly approved — last {window_days} days",
        f"Approved spend: ${total:.0f} · Potential save: ${saveable:.0f} · Items: {len(rows)}",
        "",
        "COST TABLE (re-sourced comps first)",
        "-" * 56,
    ]
    for r in rows:
        item = r["item"]
        better = r.get("better")
        verdict = report_verdict(r)
        price = float(item.get("price") or 0)
        lines.append(
            f"[{verdict}] ${price:.0f}  {item.get('brand', '')} — {(item.get('title') or '')[:70]}"
        )
        lines.append(f"  yours: {item.get('url', '')}")
        if better:
            lines.append(
                f"  better: ${float(better.get('price') or 0):.0f} "
                f"({better.get('why')}) → {better.get('url', '')}"
            )
        elif r.get("market_low") is not None:
            lines.append(f"  market low seen: ${r['market_low']:.0f}")
        else:
            lines.append("  comps: no stronger alt found")
        lines.append("")
    return "\n".join(lines)


def send_report_email(subject: str, text_body: str, html_body: str) -> bool:
    if not (REPORT_EMAIL and SMTP_USER and SMTP_PASSWORD):
        log.warning("Email report skipped — set REPORT_EMAIL, SMTP_USER, SMTP_PASSWORD in .env")
        return False
    msg = MIMEMultipart("alternative")
    msg["Subject"] = subject
    msg["From"] = REPORT_FROM or SMTP_USER
    msg["To"] = REPORT_EMAIL
    msg.attach(MIMEText(text_body, "plain", "utf-8"))
    msg.attach(MIMEText(html_body, "html", "utf-8"))
    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as server:
        server.starttls()
        server.login(SMTP_USER, SMTP_PASSWORD)
        server.sendmail(msg["From"], [REPORT_EMAIL], msg.as_string())
    log.info("Weekly report emailed to %s", REPORT_EMAIL)
    return True


async def run_weekly_report(app=None, days: int = WEEKLY_REPORT_LOOKBACK_DAYS) -> Path | None:
    """Scour comps for approved items, then email + Telegram a cost-first report."""
    items = approved_in_lookback(days)
    log.info("Weekly report: scouring comps for %s approved items...", len(items))
    rows = await scour_approved_comps(items)
    # Sort table by your cost (highest first) after comps are attached
    rows.sort(key=lambda r: float(r["item"].get("price") or 0), reverse=True)

    html_body = build_weekly_report_html(rows, days)
    text_body = build_weekly_report_text(rows, days)

    REPORTS_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    html_path = REPORTS_DIR / f"weekly_approved_{stamp}.html"
    txt_path = REPORTS_DIR / f"weekly_approved_{stamp}.txt"
    html_path.write_text(html_body, encoding="utf-8")
    txt_path.write_text(text_body, encoding="utf-8")

    subject = f"feling. weekly approved — {datetime.now().strftime('%Y-%m-%d')}"
    emailed = send_report_email(subject, text_body, html_body)

    if app is not None and TELEGRAM_CHAT_ID:
        total = sum(float(r["item"].get("price") or 0) for r in rows)
        saveable = sum(float((r.get("better") or {}).get("savings") or 0) for r in rows)
        switches = [r for r in rows if report_verdict(r) in {"SWITCH", "UPGRADE"}][:5]
        summary = (
            f"📊 Weekly approved report\n"
            f"Items: {len(rows)} · Spend: ${total:.0f} · Potential save: ${saveable:.0f}\n"
            f"Email: {'sent to ' + REPORT_EMAIL if emailed else 'not configured (saved locally)'}\n"
            f"File: {html_path.name}\n"
        )
        if switches:
            summary += "\nTop SWITCH / UPGRADE calls:\n"
            for r in switches:
                item = r["item"]
                better = r["better"]
                summary += (
                    f"• ${float(item.get('price') or 0):.0f} → "
                    f"${float(better.get('price') or 0):.0f} "
                    f"{(item.get('title') or '')[:40]}\n"
                )
        await app.bot.send_message(chat_id=TELEGRAM_CHAT_ID, text=summary[:3500])
        try:
            with html_path.open("rb") as f:
                await app.bot.send_document(
                    chat_id=TELEGRAM_CHAT_ID,
                    document=InputFile(f, filename=html_path.name),
                    caption="Open in browser for the cost table",
                )
        except Exception as e:
            log.warning("Could not send report document: %s", e)

    log.info("Weekly report written to %s (email=%s)", html_path, emailed)
    return html_path


# ─── Weekly procurement export (newly approved → Excel) ───────────────────────

def load_last_procurement_export() -> datetime | None:
    data = load_json(EXPORT_MARKER_DB, {})
    raw = data.get("last_export_at")
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw)
    except ValueError:
        return None


def save_last_procurement_export(ts: datetime):
    save_json(EXPORT_MARKER_DB, {"last_export_at": ts.isoformat()})


def newly_approved_for_procurement() -> list[dict]:
    """Items approved since the last procurement export (first run: last N days)."""
    approved = load_json(APPROVED_DB, [])
    cutoff = load_last_procurement_export() or (
        datetime.now() - timedelta(days=PROCUREMENT_EXPORT_LOOKBACK_DAYS)
    )
    fresh = []
    for item in approved:
        raw = item.get("approved_at") or ""
        try:
            ts = datetime.fromisoformat(raw)
        except ValueError:
            continue
        if ts > cutoff:
            fresh.append(item)
    fresh.sort(key=lambda x: x.get("approved_at") or "")
    return fresh


def build_procurement_workbook(items: list[dict]) -> Path:
    """Build a procurement-ready .xlsx of newly approved items."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    tracker = load_json(INVENTORY_STATUS_DB, {})

    wb = Workbook()
    ws = wb.active
    ws.title = "Newly Approved"

    headers = [
        "Brand", "Item", "Platform", "Region", "Condition",
        "Price (USD)", "Source Link", "Photo", "Approved Date",
        "Order Status", "Notes",
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1A1A1A")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
        cell.fill = header_fill

    for item in items:
        approved_at = (item.get("approved_at") or "")[:10]
        tracked = tracker.get(listing_uid(item), {})
        ws.append([
            item.get("brand", ""),
            item.get("title", ""),
            item.get("platform", ""),
            item.get("region_label") or item.get("region") or "",
            item.get("condition", ""),
            float(item.get("price") or 0),
            item.get("url", ""),
            item.get("photo", ""),
            approved_at,
            STATUS_LABELS.get(tracked.get("status"), "Approved (not yet purchased)"),
            tracked.get("notes", ""),
        ])
        r = ws.max_row
        if item.get("url"):
            link_cell = ws.cell(row=r, column=7)
            link_cell.hyperlink = item["url"]
            link_cell.style = "Hyperlink"
        ws.cell(row=r, column=6).number_format = "$#,##0.00"

    widths = [20, 45, 12, 14, 14, 12, 42, 42, 14, 14, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    REPORTS_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d")
    path = REPORTS_DIR / f"procurement_newly_approved_{stamp}.xlsx"
    wb.save(path)
    return path


def send_procurement_email(path: Path, count: int) -> bool:
    if not (REPORT_EMAIL and SMTP_USER and SMTP_PASSWORD):
        log.warning("Procurement email skipped — set REPORT_EMAIL, SMTP_USER, SMTP_PASSWORD in .env")
        return False
    msg = MIMEMultipart()
    msg["Subject"] = (
        f"feling. procurement sheet — {count} new approval"
        f"{'s' if count != 1 else ''} ({datetime.now().strftime('%Y-%m-%d')})"
    )
    msg["From"] = REPORT_FROM or SMTP_USER
    msg["To"] = REPORT_EMAIL
    msg.attach(MIMEText(
        f"Attached: {count} newly approved item(s) since the last export, ready for procurement.",
        "plain", "utf-8",
    ))
    with path.open("rb") as f:
        part = MIMEApplication(f.read(), _subtype="xlsx")
    part.add_header("Content-Disposition", "attachment", filename=path.name)
    msg.attach(part)
    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as server:
        server.starttls()
        server.login(SMTP_USER, SMTP_PASSWORD)
        server.sendmail(msg["From"], [REPORT_EMAIL], msg.as_string())
    log.info("Procurement sheet emailed to %s (%s items)", REPORT_EMAIL, count)
    return True


async def run_procurement_export(app=None) -> Path | None:
    """Compile newly approved items since the last export into a procurement Excel sheet."""
    items = newly_approved_for_procurement()
    now = datetime.now()

    if not items:
        log.info("Procurement export: no new approvals since last export.")
        if app is not None and TELEGRAM_CHAT_ID:
            await app.bot.send_message(
                chat_id=TELEGRAM_CHAT_ID,
                text="📋 Procurement export: no newly approved items since last time.",
            )
        save_last_procurement_export(now)
        return None

    path = build_procurement_workbook(items)
    emailed = send_procurement_email(path, len(items))
    save_last_procurement_export(now)

    if app is not None and TELEGRAM_CHAT_ID:
        summary = (
            f"📋 Procurement export ready — {len(items)} newly approved item(s)\n"
            f"Email: {'sent to ' + REPORT_EMAIL if emailed else 'not configured (saved locally)'}\n"
            f"File: {path.name}"
        )
        await app.bot.send_message(chat_id=TELEGRAM_CHAT_ID, text=summary)
        try:
            with path.open("rb") as f:
                await app.bot.send_document(
                    chat_id=TELEGRAM_CHAT_ID,
                    document=InputFile(f, filename=path.name),
                    caption="Newly approved — for procurement",
                )
        except Exception as e:
            log.warning("Could not send procurement document: %s", e)

    log.info(
        "Procurement export written to %s (%s items, email=%s)",
        path, len(items), emailed,
    )
    return path


# ─── Shared approve/skip core (used by Telegram buttons AND the team web API) ──

async def _notify_telegram(text: str):
    """Best-effort notification to the owner's Telegram chat from non-Telegram callers."""
    if _BOT_REF is None or not TELEGRAM_CHAT_ID:
        return
    try:
        await _BOT_REF.send_message(chat_id=TELEGRAM_CHAT_ID, text=text[:3500])
    except Exception as e:
        log.warning("Telegram notify failed: %s", e)


async def approve_pending_item(uid: str, item: dict, actor: str) -> dict:
    """Core approve logic. `actor` is a human-readable source, e.g. "telegram" or
    "web:jane". Runs the website publish step in the background so callers don't
    block on it, and notifies Telegram either way so all approvals stay visible
    in one place regardless of which surface (Telegram or the web /review queue)
    triggered them.
    """
    pending = load_json(PENDING_DB, {})
    item = normalize_item_brand(item)
    item["approved_at"] = datetime.now().isoformat()
    item["approved_by"] = actor
    approved = upsert_approved(item)
    save_json(APPROVED_DB, approved)

    pending.pop(uid, None)
    save_json(PENDING_DB, pending)

    profile = rebuild_taste_profile()
    category = categorize_item(item)

    async def _publish_later():
        try:
            status = await asyncio.to_thread(publish_to_website, item)
            await _notify_telegram(f"🌐 Website sync: {status}")
        except Exception as e:
            log.warning("Background website publish failed: %s", e)

    asyncio.create_task(_publish_later())

    if actor != "telegram":
        await _notify_telegram(
            f"✅ Approved via web by {actor.removeprefix('web:')}\n"
            f"{item.get('brand', '')} — {item.get('title', '')[:80]}\n"
            f"${float(item.get('price') or 0):.0f} • {item.get('platform', '')} • {category}"
        )

    return {
        "item": item,
        "category": category,
        "taste_top": profile.get("positive_tokens", [])[:5],
        "approved_count": profile.get("approved_count", 0),
    }


async def skip_pending_item(uid: str, item: dict, actor: str) -> dict:
    """Core skip logic, shared the same way as approve_pending_item."""
    pending = load_json(PENDING_DB, {})
    skipped = load_json(SKIPPED_DB, [])
    item = dict(item)
    item["skipped_at"] = datetime.now().isoformat()
    item["skipped_by"] = actor
    skipped.append(item)
    save_json(SKIPPED_DB, skipped[-200:])

    pending.pop(uid, None)
    save_json(PENDING_DB, pending)
    rebuild_taste_profile()

    if actor != "telegram":
        await _notify_telegram(
            f"❌ Skipped via web by {actor.removeprefix('web:')}\n"
            f"{item.get('brand', '')} — {item.get('title', '')[:80]}"
        )

    return {"item": item}


# ─── Team web API (auth-gated: /ops and /review pages talk to this) ───────────

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_user(username: str, password: str) -> bool:
    users = load_json(USERS_DB, {})
    record = users.get(username)
    if not record:
        return False
    try:
        return bcrypt.checkpw(password.encode("utf-8"), record["password_hash"].encode("utf-8"))
    except Exception:
        return False


def add_user_cli():
    """`python agent.py --add-user` — create or reset a named team login."""
    import getpass
    username = input("Username: ").strip()
    if not username:
        print("Username required.")
        return
    password = getpass.getpass("Password (min 8 chars): ")
    if len(password) < 8:
        print("Password must be at least 8 characters.")
        return
    users = load_json(USERS_DB, {})
    users[username] = {
        "password_hash": hash_password(password),
        "created_at": datetime.now().isoformat(),
    }
    save_json(USERS_DB, users)
    print(f"User '{username}' saved to {USERS_DB}.")


def bootstrap_user_from_env():
    """Non-interactive alternative to --add-user for hosts with no shell access
    (e.g. Railway). Set BOOTSTRAP_USER / BOOTSTRAP_PASSWORD once, redeploy, then
    remove them — safe to leave the vars in place since this no-ops once that
    username already exists.
    """
    username = (os.getenv("BOOTSTRAP_USER") or "").strip()
    password = os.getenv("BOOTSTRAP_PASSWORD") or ""
    if not username or not password:
        return
    users = load_json(USERS_DB, {})
    if username in users:
        return
    users[username] = {
        "password_hash": hash_password(password),
        "created_at": datetime.now().isoformat(),
    }
    save_json(USERS_DB, users)
    log.info("Bootstrapped team account '%s' from BOOTSTRAP_USER/BOOTSTRAP_PASSWORD.", username)


def create_session(username: str) -> str:
    token = secrets.token_urlsafe(32)
    _SESSIONS[token] = {
        "username": username,
        "expires": datetime.now() + timedelta(hours=SESSION_TTL_HOURS),
    }
    return token


def session_username(request: web.Request) -> str | None:
    token = request.cookies.get("session_id")
    if not token:
        return None
    session = _SESSIONS.get(token)
    if not session:
        return None
    if datetime.now() > session["expires"]:
        _SESSIONS.pop(token, None)
        return None
    return session["username"]


@web.middleware
async def auth_middleware(request: web.Request, handler):
    if request.path == "/api/auth/login":
        return await handler(request)
    username = session_username(request)
    if not username:
        return web.json_response({"error": "Not signed in"}, status=401)
    request["username"] = username
    return await handler(request)


async def api_login(request: web.Request):
    try:
        data = await request.json()
    except Exception:
        return web.json_response({"error": "Invalid JSON body"}, status=400)
    username = str(data.get("username") or "").strip()
    password = str(data.get("password") or "")
    if not verify_user(username, password):
        return web.json_response({"error": "Invalid username or password"}, status=401)
    token = create_session(username)
    resp = web.json_response({"username": username})
    resp.set_cookie(
        "session_id", token,
        httponly=True, samesite="Lax", max_age=SESSION_TTL_HOURS * 3600,
    )
    return resp


async def api_logout(request: web.Request):
    token = request.cookies.get("session_id")
    if token:
        _SESSIONS.pop(token, None)
    resp = web.json_response({"ok": True})
    resp.del_cookie("session_id")
    return resp


async def api_me(request: web.Request):
    return web.json_response({"username": request["username"]})


async def api_pending_list(request: web.Request):
    pending = load_json(PENDING_DB, {})
    return web.json_response([{**item, "uid": uid} for uid, item in pending.items()])


async def api_pending_approve(request: web.Request):
    uid = request.match_info["uid"]
    pending = load_json(PENDING_DB, {})
    item = pending.get(uid)
    if not item:
        return web.json_response({"error": "Item not found or already processed"}, status=404)
    result = await approve_pending_item(uid, item, actor=f"web:{request['username']}")
    return web.json_response(result["item"])


async def api_pending_skip(request: web.Request):
    uid = request.match_info["uid"]
    pending = load_json(PENDING_DB, {})
    item = pending.get(uid)
    if not item:
        return web.json_response({"error": "Item not found or already processed"}, status=404)
    result = await skip_pending_item(uid, item, actor=f"web:{request['username']}")
    return web.json_response(result["item"])


async def api_inventory_list(request: web.Request):
    approved = load_json(APPROVED_DB, [])
    tracker = load_json(INVENTORY_STATUS_DB, {})
    merged = []
    for item in approved:
        uid = listing_uid(item)
        tracked = tracker.get(uid, {})
        merged.append({
            "uid": uid,
            "platform": item.get("platform"),
            "brand": item.get("brand"),
            "title": item.get("title"),
            "cost": float(item.get("price") or 0),
            "condition": item.get("condition"),
            "url": item.get("url"),
            "photo": item.get("photo"),
            "approved_at": item.get("approved_at"),
            "status": tracked.get("status", "approved"),
            **tracked,
        })
    return web.json_response(merged)


async def api_inventory_patch(request: web.Request):
    uid = request.match_info["uid"]
    try:
        patch = await request.json()
    except Exception:
        return web.json_response({"error": "Invalid JSON body"}, status=400)
    tracker = load_json(INVENTORY_STATUS_DB, {})
    updated = {
        **tracker.get(uid, {}),
        **patch,
        "updated_at": datetime.now().isoformat(),
        "updated_by": request["username"],
    }
    tracker[uid] = updated
    save_json(INVENTORY_STATUS_DB, tracker)
    return web.json_response({"uid": uid, **updated})


async def suggest_resale_price(item: dict) -> dict:
    """Sell-side comp search: what similar pieces are actively asking right now,
    as opposed to the buy-side weekly report, which re-sources cheaper
    alternatives for what you *paid*. Reuses the same Grailed comp
    infrastructure (build_comp_query / search_grailed_comps / title_similarity)
    that already powers /report.
    """
    query = build_comp_query(item)
    async with AsyncSession(impersonate="chrome", headers=HEADERS) as session:
        comps = await search_grailed_comps(session, query, limit=30)

    own_url = (item.get("url") or "").rstrip("/")
    similar = [
        c for c in comps
        if float(c.get("price") or 0) > 0
        and (c.get("url") or "").rstrip("/") != own_url
        and title_similarity(item.get("title", ""), c.get("title", "")) >= 0.2
    ]
    prices = sorted(float(c["price"]) for c in similar)

    suggestion = None
    if prices:
        n = len(prices)
        median = prices[n // 2] if n % 2 else (prices[n // 2 - 1] + prices[n // 2]) / 2
        suggestion = {
            "median": round(median, 0),
            "low": round(prices[0], 0),
            "high": round(prices[-1], 0),
            "sample_size": n,
        }

    similar.sort(key=lambda c: float(c.get("price") or 0))
    return {"query": query, "suggestion": suggestion, "comps": similar[:8]}


async def api_inventory_price_suggestion(request: web.Request):
    uid = request.match_info["uid"]
    approved = load_json(APPROVED_DB, [])
    item = next((i for i in approved if listing_uid(i) == uid), None)
    if not item:
        return web.json_response({"error": "Item not found"}, status=404)
    try:
        result = await suggest_resale_price(item)
    except Exception as e:
        log.warning("Price suggestion failed for %s: %s", uid, e)
        return web.json_response({"error": "Comp search failed — try again"}, status=502)
    return web.json_response(result)


def build_web_app() -> web.Application:
    app = web.Application(middlewares=[auth_middleware])
    app.router.add_post("/api/auth/login", api_login)
    app.router.add_post("/api/auth/logout", api_logout)
    app.router.add_get("/api/auth/me", api_me)
    app.router.add_get("/api/pending", api_pending_list)
    app.router.add_post("/api/pending/{uid}/approve", api_pending_approve)
    app.router.add_post("/api/pending/{uid}/skip", api_pending_skip)
    app.router.add_get("/api/inventory", api_inventory_list)
    app.router.add_patch("/api/inventory/{uid}", api_inventory_patch)
    app.router.add_get("/api/inventory/{uid}/price-suggestion", api_inventory_price_suggestion)
    return app


# ─── Telegram Handlers ────────────────────────────────────────────────────────

async def cmd_start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "👋 *feling. Sourcing Agent* is live.\n\n"
        "I'm monitoring Grailed, Vestiaire, and Yahoo Auctions Japan (via Buyee checkout).\n"
        "Taste lane: Chanel/Dior rare bags, Dior shoes & clothes, Cavalli, Galliano, Tom Ford Gucci, Gaultier.\n"
        "Inspiration: Isle of Monday / Jean Vintage / Break Archive.\n"
        "Japan lane is cheaper by default.\n\n"
        "When I find something that matches your criteria I'll send it here with photo, price, "
        "condition, and platform. Tap *Approve* to buy or *Skip* to dismiss.\n\n"
        "Commands:\n"
        "/start — this message\n"
        "/scan — trigger a manual scan now\n"
        "/price — show or set global price filter\n"
        "/trends — demand analysis (your taste + market heat)\n"
        "/taste — show learned aesthetic from Approves\n"
        "/rules — show current sourcing rules\n"
        "/pending — show items awaiting decision\n"
        "/approved — reopen all saved buys with Open buttons\n"
        "/report — weekly cost report (re-sources comps, emails you)\n"
        "/procurement — export newly approved items to Excel (auto every Friday)",
        parse_mode="Markdown"
    )


async def cmd_scan(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("🔍 Running manual scan now...")
    await run_sourcing_scan(ctx.application)
    await update.message.reply_text("✅ Scan complete.")


async def cmd_report(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Manual weekly-style report. Optional: /report 14 for lookback days."""
    days = WEEKLY_REPORT_LOOKBACK_DAYS
    if ctx.args:
        try:
            days = max(1, min(90, int(ctx.args[0])))
        except ValueError:
            pass
    await update.message.reply_text(
        f"📊 Building approved cost report (last {days} days).\n"
        "Re-sourcing comps on Grailed first — this can take a minute..."
    )
    try:
        path = await run_weekly_report(ctx.application, days=days)
        await update.message.reply_text(
            f"✅ Report ready.\n"
            f"Saved: {path}\n"
            f"Email: {'sent' if REPORT_EMAIL and SMTP_USER and SMTP_PASSWORD else 'add REPORT_EMAIL + SMTP_* to .env'}"
        )
    except Exception as e:
        log.exception("Weekly report failed")
        await update.message.reply_text(f"Report failed: {e}")


async def cmd_procurement(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Manual trigger: compile newly approved items into an Excel sheet for procurement."""
    await update.message.reply_text("📋 Compiling newly approved items into an Excel sheet...")
    try:
        path = await run_procurement_export(ctx.application)
        if path is None:
            await update.message.reply_text("No newly approved items since the last export.")
        else:
            await update.message.reply_text(f"✅ Procurement sheet ready: {path}")
    except Exception as e:
        log.exception("Procurement export failed")
        await update.message.reply_text(f"Procurement export failed: {e}")


async def cmd_price(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Set or view price filters.
    /price
    /price 100 600
    /price clear
    /price japan 300
    /price japan clear
    """
    settings = load_settings()
    args = ctx.args or []

    if not args:
        gmin = settings.get("global_min_price")
        gmax = settings.get("global_max_price")
        jmax = settings.get("japan_max_price")
        lines = []
        if gmin is None and gmax is None:
            lines.append("No global price filter.")
        else:
            lines.append(f"Global filter: ${gmin or 0:.0f}–${gmax if gmax is not None else '∞'}")
        lines.append(
            f"Japan cheaper ceiling: ${jmax:.0f}" if jmax is not None
            else f"Japan cheaper ceiling: default ${DEFAULT_JAPAN_MAX_USD} (per brand jp_max)"
        )
        lines.append(
            "\nUsage:\n"
            "/price 100 600 — global $100–$600\n"
            "/price clear — remove global filter\n"
            "/price japan 300 — Japan auctions max $300\n"
            "/price japan clear — reset Japan ceiling"
        )
        await update.message.reply_text("\n".join(lines))
        return

    if args[0].lower() == "japan":
        if len(args) == 1:
            jmax = settings.get("japan_max_price")
            await update.message.reply_text(
                f"Japan ceiling: ${jmax:.0f}" if jmax is not None
                else f"Japan ceiling: default ${DEFAULT_JAPAN_MAX_USD}"
            )
            return
        if args[1].lower() in {"clear", "reset", "off"}:
            settings["japan_max_price"] = None
            save_settings(settings)
            await update.message.reply_text("Japan ceiling reset to brand defaults.")
            return
        try:
            jmax = float(args[1].replace("$", "").replace(",", ""))
        except ValueError:
            await update.message.reply_text("Usage: /price japan 300")
            return
        settings["japan_max_price"] = jmax
        save_settings(settings)
        await update.message.reply_text(f"✅ Japan cheaper ceiling set: max ${jmax:.0f}")
        return

    if len(args) == 1 and args[0].lower() in {"clear", "reset", "off"}:
        settings["global_min_price"] = None
        settings["global_max_price"] = None
        save_settings(settings)
        await update.message.reply_text("Global price filter cleared.")
        return

    if len(args) != 2:
        await update.message.reply_text("Usage: /price 100 600   or   /price japan 300")
        return

    try:
        lo = float(args[0].replace("$", "").replace(",", ""))
        hi = float(args[1].replace("$", "").replace(",", ""))
    except ValueError:
        await update.message.reply_text("Prices must be numbers. Example: /price 150 700")
        return

    if lo > hi:
        lo, hi = hi, lo

    settings["global_min_price"] = lo
    settings["global_max_price"] = hi
    save_settings(settings)
    await update.message.reply_text(
        f"✅ Global price filter set: ${lo:.0f}–${hi:.0f}\n"
        "Next /scan will only send items in this range."
    )


async def cmd_trends(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("📈 Analyzing your approvals + live Grailed heat...")
    try:
        report = await build_trend_report()
        await update.message.reply_text(report)
    except Exception as e:
        log.error("Trend report failed: %s", e)
        await update.message.reply_text(f"Trend analysis failed: {e}")


async def cmd_taste(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    try:
        args = [a.lower() for a in (ctx.args or [])]
        if args and args[0] in {"retrain", "refresh", "rebuild"}:
            profile = rebuild_taste_profile()
            await update.message.reply_text("🧠 Taste profile retrained from your Approves/Skips.")
        else:
            profile = load_taste_profile()

        if not profile.get("approved_count"):
            await update.message.reply_text(
                "No training data yet.\n"
                f"Approve at least {TASTE_MIN_SAMPLES} pieces and the bot will start refining."
            )
            return

        ready = "ON" if profile.get("ready") else f"warming up ({profile.get('approved_count')}/{TASTE_MIN_SAMPLES})"
        pos = ", ".join(profile.get("positive_tokens", [])[:10]) or "—"
        neg = ", ".join(profile.get("negative_tokens", [])[:6]) or "—"
        queries = "\n".join(f"• {q}" for q in profile.get("learned_queries", [])[:6]) or "• none yet"
        avg = profile.get("avg_approved_price")
        avg_line = f"${avg:.0f}" if avg is not None else "n/a"

        await update.message.reply_text(
            f"🧠 Learned aesthetic\n"
            f"Status: {ready}\n"
            f"Approves: {profile.get('approved_count', 0)} • Skips: {profile.get('skipped_count', 0)}\n"
            f"Avg approved price: {avg_line}\n\n"
            f"North star: Chanel/Dior rare bags • Dior shoes/clothes • Cavalli • Galliano • Tom Ford Gucci • Gaultier\n"
            f"Inspiration: {', '.join(profile.get('inspiration', [])[:3]) or 'Isle of Monday / Jean Vintage / Break Archive'}\n\n"
            f"Boosting: {pos}\n"
            f"Downranking: {neg}\n\n"
            f"Learned search queries:\n{queries}\n\n"
            f"Tip: keep Approving/Skipping — each tap retrains the filter."
        )
    except Exception as e:
        log.exception("cmd_taste failed")
        if update.message:
            await update.message.reply_text(f"Taste command failed: {e}")


async def cmd_rules(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    settings = load_settings()
    lines = ["*Current sourcing rules:*\n"]
    if settings.get("global_min_price") is not None or settings.get("global_max_price") is not None:
        lines.append(
            f"Global price filter: ${settings.get('global_min_price') or 0:.0f}"
            f"–${settings.get('global_max_price') if settings.get('global_max_price') is not None else '∞'}\n"
        )
    for r in RULES:
        lo, hi = effective_price_bounds(r)
        jlo, jhi = effective_price_bounds(r, platform="YahooJP")
        lines.append(
            f"🏷 *{r['brand']}*\n"
            f"  Focus: {r.get('focus', 'general')} — {', '.join((r.get('search_queries') or r['keywords'])[:3])}\n"
            f"  JP: {', '.join((r.get('jp_queries') or [])[:2]) or '—'}\n"
            f"  CN/TW: {', '.join((r.get('tw_queries') or r.get('cn_queries') or [])[:2]) or '—'}\n"
            f"  US/EU: ${lo:.0f}–${hi:.0f}\n"
            f"  Japan: ${jlo:.0f}–${jhi:.0f}\n"
            f"  Conditions: {', '.join(r['conditions'][:3])}...\n"
        )
    await update.message.reply_text("\n".join(lines), parse_mode="Markdown")


async def cmd_pending(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    pending = load_json(PENDING_DB, {})
    if not pending:
        await update.message.reply_text("No pending items.")
        return
    lines = [f"*{len(pending)} pending items:*\n"]
    for uid, item in list(pending.items())[:10]:
        lines.append(f"• {item['brand']} — {item['title'][:40]} — ${item['price']:.0f} ({item['platform']})")
    await update.message.reply_text("\n".join(lines), parse_mode="Markdown")


async def cmd_approved(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    approved = load_json(APPROVED_DB, [])
    if not approved:
        await update.message.reply_text("No approved items yet.")
        return

    await update.message.reply_text(
        f"✅ {len(approved)} approved items — tap Open to buy (opens in browser):"
    )
    for item in approved[-10:]:
        text = (
            f"✅ {item['brand']}\n"
            f"{item['title'][:80]}\n"
            f"${item['price']:.0f} • {item['platform']} • {item.get('approved_at', '')[:10]}"
        )
        keyboard = InlineKeyboardMarkup([
            [InlineKeyboardButton("🔗 Open listing", url=item["url"])]
        ]) if item.get("url") else None
        await update.message.reply_text(text, reply_markup=keyboard)


async def handle_button(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """Handle Approve / Skip button taps."""
    query   = update.callback_query
    await query.answer()

    data = query.data or ""

    if data == "show_approved":
        approved = load_json(APPROVED_DB, [])
        if not approved:
            await query.message.reply_text("No approved items yet.")
            return
        await query.message.reply_text(
            f"✅ {len(approved)} approved — latest first:"
        )
        for row in approved[-8:]:
            text = (
                f"✅ {row['brand']}\n"
                f"{row['title'][:80]}\n"
                f"${row['price']:.0f} • {row['platform']}"
            )
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton("🔗 Open listing", url=row["url"])]
            ]) if row.get("url") else None
            await query.message.reply_text(text, reply_markup=kb)
        return

    try:
        action, uid = data.split(":", 1)
    except ValueError:
        await query.answer("Bad button data.", show_alert=True)
        return

    pending = load_json(PENDING_DB, {})
    item    = pending.get(uid)

    if not item:
        try:
            await append_message_status(query, "\n\nAlready processed.")
        except Exception as e:
            log.warning("Could not update already-processed message: %s", e)
            await query.message.reply_text("Already processed.")
        return

    if action == "approve":
        result = await approve_pending_item(uid, item, actor="telegram")
        item = result["item"]
        category = result["category"]
        top = ", ".join(result["taste_top"]) or "still learning"

        open_keyboard = InlineKeyboardMarkup([
            [InlineKeyboardButton("🔗 Open listing to buy", url=item["url"])],
            [InlineKeyboardButton("📋 View all approved", callback_data="show_approved")],
        ])

        try:
            await append_message_status(query, "\n\n✅ Approved")
        except Exception as e:
            log.warning("Could not edit approved caption: %s", e)

        await query.message.reply_text(
            f"✅ Saved to approved\n"
            f"{item['brand']} — {item['title'][:80]}\n"
            f"${item['price']:.0f} • {item['platform']} • {category}\n\n"
            f"🧠 Taste updated ({result['approved_count']} saves)\n"
            f"Learning: {top}\n\n"
            f"Tap Open to buy.",
            reply_markup=open_keyboard,
        )

    elif action == "skip":
        await skip_pending_item(uid, item, actor="telegram")
        try:
            await append_message_status(query, "\n\n❌ Skipped.")
        except Exception as e:
            log.warning("Could not edit skipped caption: %s", e)
            await query.message.reply_text("❌ Skipped.")

# ─── Entrypoint ───────────────────────────────────────────────────────────────

async def main():
    global _BOT_REF

    if (
        not TELEGRAM_TOKEN
        or not TELEGRAM_CHAT_ID
        or TELEGRAM_TOKEN == "your_bot_token_here"
        or TELEGRAM_CHAT_ID == "your_chat_id_here"
    ):
        raise RuntimeError("Set real TELEGRAM_TOKEN and TELEGRAM_CHAT_ID values in .env")

    bot_id = TELEGRAM_TOKEN.split(":", 1)[0]
    if str(TELEGRAM_CHAT_ID).strip() == bot_id:
        raise RuntimeError(
            "TELEGRAM_CHAT_ID is currently set to the bot's own ID. "
            "Message @userinfobot on Telegram, copy YOUR personal numeric chat ID, "
            "and put that in .env instead."
        )

    bootstrap_user_from_env()
    if not load_json(USERS_DB, {}):
        log.warning(
            "No team accounts in %s yet — the /ops and /review web pages won't be "
            "reachable until you run `python agent.py --add-user`, or set "
            "BOOTSTRAP_USER / BOOTSTRAP_PASSWORD and redeploy.",
            USERS_DB,
        )

    async def on_error(update: object, context: ContextTypes.DEFAULT_TYPE):
        log.error("Telegram handler error: %s", context.error, exc_info=context.error)

    async def post_init(application: Application):
        await application.bot.set_my_commands([
            BotCommand("start", "Show help"),
            BotCommand("scan", "Run a scan now"),
            BotCommand("taste", "Show learned aesthetic"),
            BotCommand("trends", "Demand / heat report"),
            BotCommand("rules", "Show sourcing rules"),
            BotCommand("price", "Set price filters"),
            BotCommand("pending", "Items awaiting decision"),
            BotCommand("approved", "Reopen saved buys"),
            BotCommand("report", "Weekly cost report + comps"),
            BotCommand("procurement", "Export newly approved to Excel"),
        ])

    app = (
        Application.builder()
        .token(TELEGRAM_TOKEN)
        .concurrent_updates(True)
        .post_init(post_init)
        .build()
    )
    _BOT_REF = app.bot
    app.add_error_handler(on_error)
    app.add_handler(CommandHandler("start",    cmd_start))
    app.add_handler(CommandHandler("scan",     cmd_scan))
    app.add_handler(CommandHandler("price",    cmd_price))
    app.add_handler(CommandHandler("trends",   cmd_trends))
    app.add_handler(CommandHandler("taste",    cmd_taste))
    app.add_handler(CommandHandler("rules",    cmd_rules))
    app.add_handler(CommandHandler("pending",  cmd_pending))
    app.add_handler(CommandHandler("approved", cmd_approved))
    app.add_handler(CommandHandler("report",      cmd_report))
    app.add_handler(CommandHandler("procurement", cmd_procurement))
    app.add_handler(CallbackQueryHandler(handle_button))

    scheduler = AsyncIOScheduler()
    # Delay first scan so Telegram commands respond immediately after restart
    scheduler.add_job(
        run_sourcing_scan,
        "interval",
        minutes=POLL_INTERVAL,
        args=[app],
        next_run_time=datetime.now() + timedelta(seconds=45),
    )
    day_map = {
        "mon": "mon", "tue": "tue", "wed": "wed", "thu": "thu",
        "fri": "fri", "sat": "sat", "sun": "sun",
        "monday": "mon", "tuesday": "tue", "wednesday": "wed",
        "thursday": "thu", "friday": "fri", "saturday": "sat", "sunday": "sun",
    }
    report_day = day_map.get(WEEKLY_REPORT_DAY, "mon")
    scheduler.add_job(
        run_weekly_report,
        "cron",
        day_of_week=report_day,
        hour=WEEKLY_REPORT_HOUR,
        minute=0,
        kwargs={"app": app},
        id="weekly_approved_report",
        replace_existing=True,
    )
    procurement_day = day_map.get(PROCUREMENT_EXPORT_DAY, "fri")
    scheduler.add_job(
        run_procurement_export,
        "cron",
        day_of_week=procurement_day,
        hour=PROCUREMENT_EXPORT_HOUR,
        minute=0,
        kwargs={"app": app},
        id="procurement_export",
        replace_existing=True,
    )
    scheduler.start()

    web_app = build_web_app()
    runner = web.AppRunner(web_app)
    await runner.setup()
    site = web.TCPSite(runner, "0.0.0.0", WEB_PORT)

    async with app:
        await app.start()
        await app.updater.start_polling(drop_pending_updates=True)
        await site.start()

        log.info(
            "feling. sourcing agent started. Scanning every %s min. "
            "Weekly report: %s %02d:00 (email=%s). "
            "Procurement export: %s %02d:00 (email=%s). "
            "Team web API: http://localhost:%s (login required, %s account(s)).",
            POLL_INTERVAL,
            report_day,
            WEEKLY_REPORT_HOUR,
            REPORT_EMAIL or "off",
            procurement_day,
            PROCUREMENT_EXPORT_HOUR,
            REPORT_EMAIL or "off",
            WEB_PORT,
            len(load_json(USERS_DB, {})),
        )

        try:
            await asyncio.Event().wait()  # run until cancelled (Ctrl+C / SIGTERM)
        finally:
            scheduler.shutdown(wait=False)
            await runner.cleanup()
            await app.updater.stop()
            await app.stop()


if __name__ == "__main__":
    import sys
    if "--add-user" in sys.argv:
        add_user_cli()
    else:
        asyncio.run(main())
